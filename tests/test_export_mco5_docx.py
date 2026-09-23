"""
Tests สำหรับ GET /clo-achievement/export/mco5-docx (Phase 2 - ดู TASK-export-mco5.md ข้อ 6) - export
ข้อมูลประกอบ มคอ.5 เป็น Word (.docx) ตามโครงแบบฟอร์ม OBE5 BRU

ครอบคลุม: โครงเอกสาร (หัวเรื่อง, header OBE5 BRU, หมวด 1-6, ไม่มีส่วนรายบุคคล), ข้อมูลที่เติมอัตโนมัติ
ตรงกับ Phase 1 (Excel) เป๊ะ (ใช้ mco5_data_service.py ชุดเดียวกัน - ไม่คำนวณซ้ำ), placeholder
"[อาจารย์ผู้สอนกรอก]" ในส่วนที่ระบบไม่มีข้อมูล, ฟอนต์ TH Sarabun New, และสิทธิ์การเข้าถึงแบบเดียวกับ
Excel export ทุกประการ (resolve_mco5_export_access() ตัวเดียวกัน)

ใช้ fixture pattern เดียวกับ tests/test_export_mco5.py
"""
from __future__ import annotations

from io import BytesIO

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.auth import get_current_user
from app.database import get_db
from app.main import app
from app.models import (
    CLO,
    AssessmentItem,
    Course,
    CourseOffering,
    Curriculum,
    Enrollment,
    ItemCLO,
    Student,
    StudentScore,
    User,
)


@pytest.fixture()
def make_client(db_session):
    def _make(user: User) -> TestClient:
        def _override_get_db():
            yield db_session

        def _override_get_current_user():
            return user

        app.dependency_overrides[get_db] = _override_get_db
        app.dependency_overrides[get_current_user] = _override_get_current_user
        return TestClient(app)

    yield _make
    app.dependency_overrides.clear()


def _make_curriculum_course_offering(db_session, instructor=None, section="1"):
    curriculum = Curriculum(name="Test Curriculum MCO5 Docx", year=2569)
    db_session.add(curriculum)
    db_session.flush()

    course = Course(
        curriculum_id=curriculum.id,
        course_code="MCO5DOCX1",
        name_th="วิชาทดสอบ มคอ.5 Word",
        name_en="MCO5 Docx Test Course",
        credit=3,
        category="วิชาแกน",
    )
    db_session.add(course)
    db_session.flush()

    offering = CourseOffering(
        course_id=course.id,
        instructor_id=instructor.id if instructor is not None else None,
        academic_year=2569,
        semester=1,
        section=section,
    )
    db_session.add(offering)
    db_session.flush()

    return curriculum, course, offering


def _enroll_student(db_session, offering, student_id, final_grade=None):
    student = Student(
        id=student_id,
        curriculum_id=offering.course.curriculum_id,
        first_name="ทดสอบ",
        last_name=student_id,
        cohort_year=69,
        current_year_level=1,
    )
    db_session.add(student)
    db_session.add(Enrollment(student_id=student.id, offering_id=offering.id, final_grade=final_grade))
    db_session.flush()
    return student


def _add_clo_with_score(db_session, *, course, offering, clo_code, admin_user_id, scores: dict):
    clo = CLO(
        course_id=course.id,
        code=clo_code,
        description=f"ทดสอบ {clo_code}",
        pass_threshold_percent=60.00,
        created_by=admin_user_id,
        domain="knowledge",
    )
    db_session.add(clo)
    db_session.flush()

    item = AssessmentItem(offering_id=offering.id, name=f"item-{clo_code}", type="quiz", total_score=100.0)
    db_session.add(item)
    db_session.flush()
    db_session.add(ItemCLO(item_id=item.id, clo_id=clo.id, weight_percent=100.00))

    for student_id, score in scores.items():
        db_session.add(StudentScore(item_id=item.id, student_id=student_id, score_obtained=score))

    return clo


def _all_paragraph_text(doc: Document) -> str:
    return "\n".join(p.text for p in doc.paragraphs)


def _all_table_cell_texts(doc: Document) -> list[str]:
    texts = []
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                texts.append(cell.text)
    return texts


class TestDocxStructure:
    def test_export_returns_docx_with_expected_sections(self, client, db_session, admin_user):
        curriculum, course, offering = _make_curriculum_course_offering(db_session)
        student = _enroll_student(db_session, offering, "MCO5DX-001", final_grade="A")
        _add_clo_with_score(
            db_session,
            course=course,
            offering=offering,
            clo_code="CLO1",
            admin_user_id=admin_user.id,
            scores={student.id: 90.0},
        )
        db_session.commit()

        resp = client.get(f"/clo-achievement/export/mco5-docx?offering_id={offering.id}")
        assert resp.status_code == 200
        assert resp.headers["content-type"] == (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
        assert "mco5_MCO5DOCX1" in resp.headers["content-disposition"]
        assert resp.headers["content-disposition"].endswith('.docx"')

        doc = Document(BytesIO(resp.content))
        text = _all_paragraph_text(doc)

        assert "รายงานผลการดำเนินการของรายวิชา" in text
        for i in range(1, 7):
            assert f"หมวดที่ {i}" in text

        header_text = "\n".join(p.text for p in doc.sections[0].header.paragraphs)
        assert "OBE5 BRU" in header_text

        # ไม่มีส่วนข้อมูลรายบุคคลใน Word เลยไม่ว่ากรณีใด (ต่างจาก Excel)
        assert "รายบุคคล" not in text
        assert student.id not in text

    def test_body_font_is_th_sarabun_new_everywhere(self, client, db_session, admin_user):
        curriculum, course, offering = _make_curriculum_course_offering(db_session)
        db_session.commit()

        resp = client.get(f"/clo-achievement/export/mco5-docx?offering_id={offering.id}")
        doc = Document(BytesIO(resp.content))

        assert doc.styles["Normal"].font.name == "TH Sarabun New"
        for p in doc.paragraphs:
            for run in p.runs:
                assert run.font.name == "TH Sarabun New"


class TestDocxDataMatchesExcel:
    def test_clo_table_matches_get_clo_achievement(self, client, db_session, admin_user):
        curriculum, course, offering = _make_curriculum_course_offering(db_session)
        s1 = _enroll_student(db_session, offering, "MCO5DX-010")
        s2 = _enroll_student(db_session, offering, "MCO5DX-011")
        _add_clo_with_score(
            db_session,
            course=course,
            offering=offering,
            clo_code="CLO1",
            admin_user_id=admin_user.id,
            scores={s1.id: 90.0, s2.id: 30.0},
        )
        db_session.commit()

        api_resp = client.get(f"/clo-achievement?offering_id={offering.id}")
        api_clo = api_resp.json()["clo_achievements"][0]

        docx_resp = client.get(f"/clo-achievement/export/mco5-docx?offering_id={offering.id}")
        doc = Document(BytesIO(docx_resp.content))

        clo_table = doc.tables[1]  # 0=course info, 1=CLO table
        header = [c.text for c in clo_table.rows[0].cells]
        assert header == ["CLOs", "กลยุทธ์การสอน", "วิธีการประเมินผล", "ผลที่เกิดกับนักศึกษา", "แนวทางพัฒนา"]

        data_row = [c.text for c in clo_table.rows[1].cells]
        assert data_row[0].startswith("CLO1")
        assert data_row[1] == ""  # กลยุทธ์การสอน เว้นว่างเสมอ
        assert str(api_clo["passed_count"]) in data_row[3]
        assert str(api_clo["failed_count"]) in data_row[3]

    def test_grade_distribution_table_totals_100(self, client, db_session, admin_user):
        curriculum, course, offering = _make_curriculum_course_offering(db_session)
        for i in range(42):
            _enroll_student(db_session, offering, f"MCO5DX-G{i:03d}", final_grade="A")
        for i in range(3):
            _enroll_student(db_session, offering, f"MCO5DX-GI{i:03d}", final_grade="I")
        db_session.commit()

        resp = client.get(f"/clo-achievement/export/mco5-docx?offering_id={offering.id}")
        doc = Document(BytesIO(resp.content))

        grade_table = doc.tables[2]  # 0=course info, 1=CLO, 2=grade distribution
        rows_by_grade = {row.cells[0].text: [c.text for c in row.cells] for row in grade_table.rows}
        assert rows_by_grade["I"][2] == "3"
        assert rows_by_grade["I"][3] == "6.67"
        assert rows_by_grade["รวม"][2] == "45"
        assert rows_by_grade["รวม"][3] == "100.00"

    def test_status_changes_with_target_rate_affects_clo_table(self, client, db_session, admin_user):
        curriculum, course, offering = _make_curriculum_course_offering(db_session)
        s1 = _enroll_student(db_session, offering, "MCO5DX-020")
        s2 = _enroll_student(db_session, offering, "MCO5DX-021")
        _add_clo_with_score(
            db_session,
            course=course,
            offering=offering,
            clo_code="CLO1",
            admin_user_id=admin_user.id,
            scores={s1.id: 90.0, s2.id: 10.0},  # 50% rate
        )
        db_session.commit()

        low_resp = client.get(
            f"/clo-achievement/export/mco5-docx?offering_id={offering.id}&target_rate=40"
        )
        high_resp = client.get(
            f"/clo-achievement/export/mco5-docx?offering_id={offering.id}&target_rate=90"
        )

        low_doc = Document(BytesIO(low_resp.content))
        high_doc = Document(BytesIO(high_resp.content))

        low_improvement = low_doc.tables[1].rows[1].cells[4].text
        high_improvement = high_doc.tables[1].rows[1].cells[4].text

        assert low_improvement == "-"  # บรรลุ -> ไม่มีข้อความเตือน
        assert "ควรระบุแนวทางปรับปรุง" in high_improvement  # ไม่บรรลุ -> มีข้อความเตือน


class TestDocxPlaceholders:
    def test_sections_without_data_show_placeholder(self, client, db_session, admin_user):
        curriculum, course, offering = _make_curriculum_course_offering(db_session)
        db_session.commit()

        resp = client.get(f"/clo-achievement/export/mco5-docx?offering_id={offering.id}")
        doc = Document(BytesIO(resp.content))

        placeholder_paragraphs = [p for p in doc.paragraphs if p.text == "[อาจารย์ผู้สอนกรอก]"]
        # หมวด 2 ข้อ 1, ข้อ 2, หมวด 3 ข้อ 5, ข้อ 6, หมวด 4, 5, 6 = 7 จุด
        assert len(placeholder_paragraphs) == 7
        for p in placeholder_paragraphs:
            run = p.runs[0]
            assert run.font.italic is True
            assert run.font.color.rgb is not None


class TestDocxAccessControl:
    def test_instructor_not_owner_returns_403(self, make_client, db_session, admin_user):
        owner = User(
            username="mco5docx-owner", password="x", first_name="เจ้าของ", last_name="วิชา", role="instructor"
        )
        other = User(
            username="mco5docx-other", password="x", first_name="คนอื่น", last_name="ไม่เกี่ยว", role="instructor"
        )
        db_session.add_all([owner, other])
        db_session.flush()
        curriculum, course, offering = _make_curriculum_course_offering(db_session, instructor=owner)
        db_session.commit()

        resp = make_client(other).get(f"/clo-achievement/export/mco5-docx?offering_id={offering.id}")
        assert resp.status_code == 403

    def test_owning_instructor_can_export(self, make_client, db_session, admin_user):
        owner = User(
            username="mco5docx-owner2", password="x", first_name="เจ้าของ", last_name="วิชา", role="instructor"
        )
        db_session.add(owner)
        db_session.flush()
        curriculum, course, offering = _make_curriculum_course_offering(db_session, instructor=owner)
        db_session.commit()

        resp = make_client(owner).get(f"/clo-achievement/export/mco5-docx?offering_id={offering.id}")
        assert resp.status_code == 200

    def test_nonexistent_offering_returns_404(self, client):
        resp = client.get("/clo-achievement/export/mco5-docx?offering_id=999999")
        assert resp.status_code == 404

    def test_other_role_can_still_export_docx(self, make_client, db_session, admin_user):
        """ต่างจาก Excel (role อื่นไม่มีชีตรายบุคคล) - Word ไม่มีส่วนรายบุคคลอยู่แล้วไม่ว่า role ไหน จึง
        export ได้ปกติ เนื้อหาเหมือนกันหมด"""
        coordinator = User(
            username="mco5docx-coordinator",
            password="x",
            first_name="ประธาน",
            last_name="หลักสูตร",
            role="coordinator",
        )
        db_session.add(coordinator)
        db_session.flush()
        curriculum, course, offering = _make_curriculum_course_offering(db_session)
        db_session.commit()

        resp = make_client(coordinator).get(
            f"/clo-achievement/export/mco5-docx?offering_id={offering.id}"
        )
        assert resp.status_code == 200
