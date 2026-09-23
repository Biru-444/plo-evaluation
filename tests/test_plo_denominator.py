"""
Tests สำหรับ TASK-plo-denominator - สถิติ PLO ระดับรุ่น/หลักสูตร (คะแนนเฉลี่ย, ร้อยละที่บรรลุ) เปลี่ยน
ตัวหารจาก "นักศึกษาทั้งหมด" เป็น "นักศึกษาที่มีข้อมูลของ PLO นั้น" (has_data=True) พร้อมแสดง coverage
คู่กันเสมอ ครอบคลุม 7 สถานการณ์ตามสเปก (ดู TASK-plo-denominator.md ข้อ 6):
  1. นักศึกษาที่มีคะแนนจริงแต่ได้ 0% -> has_data=True, นับเข้าตัวหาร, ไม่บรรลุ
  2. นักศึกษาที่ไม่มีคะแนน CLO ของ PLO นั้นเลย -> has_data=False, ไม่เข้าตัวหาร
  3. รุ่น 10 คน มีข้อมูล 4 คน บรรลุ 3 คน -> rate = 75.0, coverage = 40.0
  4. PLO ที่ไม่มีใครมีข้อมูล -> average/rate = null ไม่ error
  5. all_plo_achieved ใช้ตัวหารเฉพาะคนที่มีข้อมูลครบทุก qualifying PLO
  6. by-year endpoint ใช้สูตรเดียวกัน
  7. export: ชีตสรุปตรงกับ cohort endpoint ใหม่, ชีตรายบุคคลแสดง "-" สำหรับไม่มีข้อมูล

ใช้ fixture pattern เดียวกับ tests/test_plo_achievement_cohort.py และ tests/test_plo_report_export.py
"""
from __future__ import annotations

from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.models import (
    CLO,
    AssessmentItem,
    CLOPLOMapping,
    Course,
    CourseOffering,
    Curriculum,
    Enrollment,
    ItemCLO,
    PLO,
    Student,
    StudentScore,
    StudyPlan,
)


def _make_curriculum_plo_course(db_session, name="Test Denominator Curriculum"):
    curriculum = Curriculum(name=name, year=2569)
    db_session.add(curriculum)
    db_session.flush()

    plo = PLO(curriculum_id=curriculum.id, code="PLO1", description_th="ทดสอบตัวหาร", category="ความรู้")
    db_session.add(plo)
    db_session.flush()

    course = Course(curriculum_id=curriculum.id, course_code="DENOM1", name_th="วิชาทดสอบตัวหาร", credit=3)
    db_session.add(course)
    db_session.flush()

    offering = CourseOffering(course_id=course.id, academic_year=2569, semester=1, section="1")
    db_session.add(offering)
    db_session.flush()

    return curriculum, plo, course, offering


def _enroll_student(db_session, curriculum, offering, student_id, cohort_year=69):
    student = Student(
        id=student_id,
        curriculum_id=curriculum.id,
        first_name="ทดสอบ",
        last_name=student_id,
        cohort_year=cohort_year,
    )
    db_session.add(student)
    db_session.add(Enrollment(student_id=student.id, offering_id=offering.id))
    db_session.flush()
    return student


def _add_clo(db_session, *, course, offering, plo, admin_user_id, clo_code="CLO1", weight_percent=100.0):
    """สร้าง CLO ผูกกับ PLO + assessment_item 1 ชิ้น (เต็ม 100, weight 100% เข้า CLO) - ยังไม่ใส่คะแนน
    (ผู้เรียกเติม StudentScore เองทีละคนเพื่อคุมว่าใครมี "ข้อมูล" บ้าง)"""
    clo = CLO(
        course_id=course.id,
        code=clo_code,
        description=f"ทดสอบ {clo_code}",
        pass_threshold_percent=60.00,
        created_by=admin_user_id,
    )
    db_session.add(clo)
    db_session.flush()
    db_session.add(CLOPLOMapping(clo_id=clo.id, plo_id=plo.id, weight_percent=weight_percent))

    item = AssessmentItem(offering_id=offering.id, name=f"item-{clo_code}", type="quiz", total_score=100.0)
    db_session.add(item)
    db_session.flush()
    db_session.add(ItemCLO(item_id=item.id, clo_id=clo.id, weight_percent=100.00))
    return clo, item


class TestHasDataDistinguishesZeroFromNoData:
    def test_student_with_real_zero_score_has_data_true_not_achieved(self, client, db_session, admin_user):
        """สถานการณ์ 1: นักศึกษาที่มีคะแนนจริงแต่ได้ 0% -> has_data=True, นับเข้าตัวหาร, ไม่บรรลุ"""
        curriculum, plo, course, offering = _make_curriculum_plo_course(db_session)
        clo, item = _add_clo(db_session, course=course, offering=offering, plo=plo, admin_user_id=admin_user.id)
        student = _enroll_student(db_session, curriculum, offering, "DENOM-001")
        db_session.add(StudentScore(item_id=item.id, student_id=student.id, score_obtained=0.0))
        db_session.commit()

        resp = client.get(f"/plo/achievement?student_id={student.id}")
        assert resp.status_code == 200
        plo_item = next(p for p in resp.json()["plo_achievements"] if p["plo_id"] == plo.id)
        assert plo_item["has_data"] is True
        assert plo_item["achieved_percent"] == 0.0
        assert plo_item["is_achieved"] is False

        cohort_resp = client.get(f"/plo/achievement/cohort?curriculum_id={curriculum.id}")
        summary = next(p for p in cohort_resp.json()["plo_summary"] if p["plo_id"] == plo.id)
        # นักศึกษาคนนี้มีข้อมูลจริง (ได้ 0% ไม่ใช่ "ไม่มีข้อมูล") -> นับเข้าตัวหาร count_with_data
        assert summary["student_count_with_data"] == 1
        assert summary["average_achieved_percent"] == 0.0
        assert summary["coverage_percent"] == 100.0

    def test_student_with_no_clo_score_has_data_false_excluded_from_denominator(
        self, client, db_session, admin_user
    ):
        """สถานการณ์ 2: นักศึกษาที่ไม่มีคะแนน CLO ของ PLO นั้นเลย -> has_data=False, ไม่เข้าตัวหาร"""
        curriculum, plo, course, offering = _make_curriculum_plo_course(db_session)
        _add_clo(db_session, course=course, offering=offering, plo=plo, admin_user_id=admin_user.id)
        student = _enroll_student(db_session, curriculum, offering, "DENOM-002")
        db_session.commit()  # ไม่มี StudentScore เลย

        resp = client.get(f"/plo/achievement?student_id={student.id}")
        plo_item = next(p for p in resp.json()["plo_achievements"] if p["plo_id"] == plo.id)
        assert plo_item["has_data"] is False
        assert plo_item["achieved_percent"] == 0.0  # ยังคงเป็น 0.0 เพื่อไม่ให้ caller เดิมพัง (ดู schema)
        assert plo_item["is_achieved"] is False

        cohort_resp = client.get(f"/plo/achievement/cohort?curriculum_id={curriculum.id}")
        summary = next(p for p in cohort_resp.json()["plo_summary"] if p["plo_id"] == plo.id)
        assert summary["student_count_with_data"] == 0
        assert summary["coverage_percent"] == 0.0
        # ไม่มีใครมีข้อมูลเลย -> average/rate ต้องเป็น None (หารไม่ได้) ไม่ใช่ 0.0
        assert summary["average_achieved_percent"] is None
        assert summary["achieved_rate_percent"] is None


class TestRateAndCoverageWithMixedData:
    def test_ten_students_four_with_data_three_achieved(self, client, db_session, admin_user):
        """สถานการณ์ 3: รุ่น 10 คน มีข้อมูล 4 คน บรรลุ 3 คน -> rate = 75.0, coverage = 40.0"""
        curriculum, plo, course, offering = _make_curriculum_plo_course(db_session)
        _clo, item = _add_clo(db_session, course=course, offering=offering, plo=plo, admin_user_id=admin_user.id)

        students = [
            _enroll_student(db_session, curriculum, offering, f"DENOM-T{i:02d}") for i in range(10)
        ]
        # 4 คนแรกมีคะแนน: 3 คนบรรลุ (>=60), 1 คนไม่บรรลุ - อีก 6 คนไม่มีคะแนนเลย (ไม่มีข้อมูล)
        db_session.add(StudentScore(item_id=item.id, student_id=students[0].id, score_obtained=90.0))
        db_session.add(StudentScore(item_id=item.id, student_id=students[1].id, score_obtained=70.0))
        db_session.add(StudentScore(item_id=item.id, student_id=students[2].id, score_obtained=65.0))
        db_session.add(StudentScore(item_id=item.id, student_id=students[3].id, score_obtained=10.0))
        db_session.commit()

        resp = client.get(f"/plo/achievement/cohort?curriculum_id={curriculum.id}")
        summary = next(p for p in resp.json()["plo_summary"] if p["plo_id"] == plo.id)

        assert summary["student_count_with_data"] == 4
        assert summary["achieved_student_count"] == 3
        assert summary["achieved_rate_percent"] == 75.0
        assert summary["coverage_percent"] == 40.0


class TestNoOneHasDataDoesNotError:
    def test_plo_with_clo_mapped_but_zero_students_with_data(self, client, db_session, admin_user):
        """สถานการณ์ 4: PLO ที่ไม่มีใครมีข้อมูล (แต่มี CLO ผูกอยู่ ไม่ใช่ "ยังไม่มี CLO ผูก") -> average/rate
        เป็น null ไม่ error"""
        curriculum, plo, course, offering = _make_curriculum_plo_course(db_session)
        _add_clo(db_session, course=course, offering=offering, plo=plo, admin_user_id=admin_user.id)
        for i in range(3):
            _enroll_student(db_session, curriculum, offering, f"DENOM-N{i:02d}")
        db_session.commit()  # ไม่มีใครมีคะแนนเลย

        resp = client.get(f"/plo/achievement/cohort?curriculum_id={curriculum.id}")
        assert resp.status_code == 200
        summary = next(p for p in resp.json()["plo_summary"] if p["plo_id"] == plo.id)
        assert summary["average_achieved_percent"] is None
        assert summary["achieved_rate_percent"] is None
        assert summary["student_count_with_data"] == 0
        assert summary["achieved_student_count"] == 0

        # export ต้องไม่ error เช่นกัน
        export_resp = client.get(f"/plo/achievement/export?curriculum_id={curriculum.id}")
        assert export_resp.status_code == 200


class TestAllPLOAchievedDenominator:
    def test_uses_students_with_complete_data_across_qualifying_plos(self, client, db_session, admin_user):
        """สถานการณ์ 5: all_plo_achieved ใช้ตัวหารเฉพาะคนที่มีข้อมูลครบทุก qualifying PLO - นักศึกษาที่
        ยังไม่มีข้อมูลของ PLO ข้อใดข้อหนึ่งไม่ควรถูกนับเป็นตัวหาร แม้จะบรรลุ PLO ข้ออื่นครบก็ตาม"""
        curriculum = Curriculum(name="Test All PLO Achieved Denominator", year=2569)
        db_session.add(curriculum)
        db_session.flush()

        plo_a = PLO(curriculum_id=curriculum.id, code="PLO-A", description_th="A", category="ความรู้")
        plo_b = PLO(curriculum_id=curriculum.id, code="PLO-B", description_th="B", category="ทักษะ")
        db_session.add_all([plo_a, plo_b])
        db_session.flush()

        course = Course(curriculum_id=curriculum.id, course_code="DENOM2", name_th="วิชาทดสอบ", credit=3)
        db_session.add(course)
        db_session.flush()
        offering = CourseOffering(course_id=course.id, academic_year=2569, semester=1, section="1")
        db_session.add(offering)
        db_session.flush()

        _clo_a, item_a = _add_clo(
            db_session, course=course, offering=offering, plo=plo_a, admin_user_id=admin_user.id, clo_code="CLO-A"
        )
        _clo_b, item_b = _add_clo(
            db_session, course=course, offering=offering, plo=plo_b, admin_user_id=admin_user.id, clo_code="CLO-B"
        )

        # s1: มีข้อมูลครบทั้ง PLO-A และ PLO-B, บรรลุทั้งคู่ -> นับเข้าทั้งตัวตั้งและตัวหารของ all_plo_achieved
        s1 = _enroll_student(db_session, curriculum, offering, "DENOM-ALL1")
        db_session.add(StudentScore(item_id=item_a.id, student_id=s1.id, score_obtained=90.0))
        db_session.add(StudentScore(item_id=item_b.id, student_id=s1.id, score_obtained=90.0))

        # s2: มีข้อมูลเฉพาะ PLO-A (บรรลุ) แต่ไม่มีข้อมูล PLO-B เลย -> ไม่ควรถูกนับเป็นตัวหารของ
        # all_plo_achieved เพราะข้อมูลไม่ครบ (ต่างจากพฤติกรรมเดิมที่จะนับเป็น "ไม่บรรลุครบ" ปนไปกับคนที่
        # สอบตกจริง)
        s2 = _enroll_student(db_session, curriculum, offering, "DENOM-ALL2")
        db_session.add(StudentScore(item_id=item_a.id, student_id=s2.id, score_obtained=90.0))

        db_session.commit()

        resp = client.get(f"/plo/achievement/cohort?curriculum_id={curriculum.id}")
        body = resp.json()
        assert body["qualifying_plo_count"] == 2
        # ตัวหารใหม่: เฉพาะ s1 ที่มีข้อมูลครบทั้ง PLO-A และ PLO-B (s2 ขาดข้อมูล PLO-B เลยไม่นับ)
        assert body["all_plo_data_complete_count"] == 1
        assert body["all_plo_achieved_count"] == 1
        assert body["all_plo_achieved_percent"] == 100.0

    def test_percent_is_none_when_no_one_has_complete_data(self, client, db_session, admin_user):
        """all_plo_data_complete_count = 0 -> all_plo_achieved_percent ต้องเป็น None ไม่ error"""
        curriculum, plo, course, offering = _make_curriculum_plo_course(db_session)
        _add_clo(db_session, course=course, offering=offering, plo=plo, admin_user_id=admin_user.id)
        _enroll_student(db_session, curriculum, offering, "DENOM-ALLN1")
        db_session.commit()  # ไม่มีใครมีคะแนนเลย

        resp = client.get(f"/plo/achievement/cohort?curriculum_id={curriculum.id}")
        body = resp.json()
        assert body["all_plo_data_complete_count"] == 0
        assert body["all_plo_achieved_percent"] is None


class TestByYearEndpointUsesSameFormula:
    def test_by_year_divides_by_students_with_data(self, client, db_session, admin_user):
        """สถานการณ์ 6: by-year endpoint ใช้สูตรเดียวกัน (หารด้วยนักศึกษาที่มีข้อมูล ไม่ใช่ทั้งหมด)"""
        curriculum, plo, course, offering = _make_curriculum_plo_course(db_session)
        _clo, item = _add_clo(db_session, course=course, offering=offering, plo=plo, admin_user_id=admin_user.id)
        db_session.add(
            StudyPlan(curriculum_id=curriculum.id, course_id=course.id, cohort_year=None, year_level=1, semester=1)
        )

        s1 = _enroll_student(db_session, curriculum, offering, "DENOM-BY1")
        s2 = _enroll_student(db_session, curriculum, offering, "DENOM-BY2")
        db_session.add(StudentScore(item_id=item.id, student_id=s1.id, score_obtained=90.0))
        # s2 ไม่มีคะแนนเลย -> ไม่มีข้อมูล
        db_session.commit()

        resp = client.get(f"/plo/achievement/by-year?curriculum_id={curriculum.id}")
        assert resp.status_code == 200
        body = resp.json()
        year1 = next(y for y in body["years"] if y["year_level"] == 1)
        summary = next(p for p in year1["plo_summary"] if p["plo_id"] == plo.id)

        assert summary["student_count_with_data"] == 1
        # หารด้วยผู้มีข้อมูล 1 คน ไม่ใช่นักศึกษาทั้งหมด 2 คน -> average/rate = 90.0/100.0 ไม่ใช่ 45.0/50.0
        assert summary["average_achieved_percent"] == 90.0
        assert summary["achieved_rate_percent"] == 100.0
        assert summary["coverage_percent"] == 50.0

    def test_by_year_none_when_no_data(self, client, db_session, admin_user):
        curriculum, plo, course, offering = _make_curriculum_plo_course(db_session)
        _add_clo(db_session, course=course, offering=offering, plo=plo, admin_user_id=admin_user.id)
        db_session.add(
            StudyPlan(curriculum_id=curriculum.id, course_id=course.id, cohort_year=None, year_level=1, semester=1)
        )
        _enroll_student(db_session, curriculum, offering, "DENOM-BY3")
        db_session.commit()

        resp = client.get(f"/plo/achievement/by-year?curriculum_id={curriculum.id}")
        body = resp.json()
        year1 = next(y for y in body["years"] if y["year_level"] == 1)
        summary = next(p for p in year1["plo_summary"] if p["plo_id"] == plo.id)
        assert summary["average_achieved_percent"] is None
        assert summary["achieved_rate_percent"] is None


class TestExportMatchesNewDenominator:
    def test_summary_sheet_matches_cohort_endpoint(self, client, db_session, admin_user):
        """สถานการณ์ 7a: ชีตสรุปตรงกับ cohort endpoint ใหม่ (ทั้งค่า None ที่แสดงเป็น "-" และค่าจริง)"""
        curriculum, plo, course, offering = _make_curriculum_plo_course(db_session)
        _clo, item = _add_clo(db_session, course=course, offering=offering, plo=plo, admin_user_id=admin_user.id)
        s1 = _enroll_student(db_session, curriculum, offering, "DENOM-EXP1")
        s2 = _enroll_student(db_session, curriculum, offering, "DENOM-EXP2")
        db_session.add(StudentScore(item_id=item.id, student_id=s1.id, score_obtained=90.0))
        # s2 ไม่มีคะแนน
        db_session.commit()

        api_resp = client.get(f"/plo/achievement/cohort?curriculum_id={curriculum.id}")
        api_summary = next(p for p in api_resp.json()["plo_summary"] if p["plo_id"] == plo.id)

        export_resp = client.get(f"/plo/achievement/export?curriculum_id={curriculum.id}")
        wb = load_workbook(BytesIO(export_resp.content))
        ws = wb["สรุปการบรรลุ PLO"]
        row = next(
            r
            for r in ws.iter_rows(min_row=2, values_only=True)
            if r[0] == "PLO1"
        )
        # คอลัมน์: PLO(0) หมวด(1) คำอธิบาย(2) จำนวนวิชาที่วัด(3) จำนวนCLO(4) นักศึกษาทั้งหมด(5) มีข้อมูล(6)
        # ความครอบคลุมข้อมูล(7) คะแนนPLOเฉลี่ย(8) บรรลุ(คน)(9) ร้อยละที่บรรลุ(10) สถานะ(11)
        assert row[6] == api_summary["student_count_with_data"] == 1
        assert row[7] == pytest.approx(api_summary["coverage_percent"], abs=0.1)
        assert row[8] == api_summary["average_achieved_percent"] == 90.0
        assert row[10] == api_summary["achieved_rate_percent"] == 100.0

    def test_personal_sheet_shows_dash_for_no_data_not_zero(self, client, db_session, admin_user):
        """สถานการณ์ 7b: ชีตรายบุคคลแสดง "-" สำหรับ has_data=False (ไม่ใช่ตัวเลข 0) และแสดงตัวเลขจริง
        (รวมถึง 0.0 ของคนที่มีคะแนนจริงแต่ได้ 0%) เป็นตัวเลข ไม่ใช่ "-\""""
        curriculum, plo, course, offering = _make_curriculum_plo_course(db_session)
        _clo, item = _add_clo(db_session, course=course, offering=offering, plo=plo, admin_user_id=admin_user.id)

        s_real_zero = _enroll_student(db_session, curriculum, offering, "DENOM-PZ1")
        s_no_data = _enroll_student(db_session, curriculum, offering, "DENOM-PZ2")
        db_session.add(StudentScore(item_id=item.id, student_id=s_real_zero.id, score_obtained=0.0))
        # s_no_data ไม่มีคะแนนเลย
        db_session.commit()

        resp = client.get(f"/plo/achievement/export?curriculum_id={curriculum.id}")
        wb = load_workbook(BytesIO(resp.content))
        ws = wb["รายบุคคล"]

        rows_by_id = {row[0]: row for row in ws.iter_rows(min_row=4, values_only=True)}
        # คอลัมน์: รหัสนักศึกษา(0) ชื่อ-สกุล(1) รุ่น(2) PLO1(3) จำนวน PLO ที่บรรลุ(4)
        assert rows_by_id["DENOM-PZ1"][3] == 0.0  # มีคะแนนจริงแต่ได้ 0% -> แสดงตัวเลข 0.0 ไม่ใช่ "-"
        assert rows_by_id["DENOM-PZ2"][3] == "-"  # ไม่มีข้อมูลเลย -> แสดง "-"
