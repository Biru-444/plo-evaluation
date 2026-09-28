"""
Tests สำหรับ GET /export/course/{course_id} และ GET /export/student/{student_id} (app/routes/export.py)
— เปิดไฟล์ .xlsx ที่ได้กลับมาด้วย openpyxl แล้วตรวจชื่อชีท + ค่าในเซลล์จริง ว่าผ่าน/ไม่ผ่าน CLO และ
ผลบรรลุ PLO ตรงกับกฎเดียวกับหน้าเว็บ (เกณฑ์ผ่าน CLO 60%, เกณฑ์บรรลุ PLO 60%) และอาจารย์ทุกคนเห็นผล CLO
ครบทุกกลุ่มเรียน (ไม่จำกัดเฉพาะกลุ่มที่ตัวเองสอน)

ทุกเทสสร้างข้อมูลของตัวเองใน db_session (rollback อัตโนมัติหลังจบเทสตาม conftest.py)
"""
from __future__ import annotations

import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.auth import get_current_user
from app.database import get_db
from app.main import app
from app.models import (
    AssessmentItem,
    CLO,
    CLOPLOMapping,
    Course,
    CourseOffering,
    Curriculum,
    Enrollment,
    ItemCLO,
    PLO,
    Student,
    StudentScore,
    StudyPlan,
    User,
    YLO,
)


@pytest.fixture()
def make_client(db_session):
    def _make(user: User) -> TestClient:
        def _override_get_db():
            yield db_session

        app.dependency_overrides[get_db] = _override_get_db
        app.dependency_overrides[get_current_user] = lambda: user
        return TestClient(app)

    yield _make
    app.dependency_overrides.clear()


def _sheet_values(ws) -> list[list[object]]:
    return [list(row) for row in ws.iter_rows(values_only=True)]


def _find_row(rows: list[list[object]], first_cell: object) -> list[object]:
    return next(row for row in rows if row and row[0] == first_cell)


def _find_row_containing(rows: list[list[object]], value: object) -> list[object]:
    return next(row for row in rows if value in row)


def _all_cells(rows: list[list[object]]) -> list[object]:
    return [cell for row in rows for cell in row]


def _make_fixtures(db_session, admin_user_id: int):
    """1 หลักสูตร, 1 วิชา (2 CLO ผูกกับ PLO1 ทั้งคู่), 1 กลุ่มเรียนของอาจารย์ "สมชาย", นักศึกษา 2 คน:
    - EXP001 ได้ 90% ทั้ง 2 CLO -> ผ่านทุก CLO -> PLO1 = 90% บรรลุ
    - EXP002 ได้ 40% ใน CLO-A, ไม่มีคะแนน CLO-B -> ไม่ผ่าน / ไม่มีข้อมูล -> PLO1 = 40% ยังไม่บรรลุ
    PLO2 ไม่มี CLO ผูกอยู่เลย -> ทุกคน "ยังไม่มีข้อมูล"
    มีอาจารย์ "อื่น" อีกคนที่ไม่ได้สอนวิชานี้ ไว้ทดสอบสิทธิ์"""
    curriculum = Curriculum(name="หลักสูตรทดสอบ Export", year=2569)
    db_session.add(curriculum)
    db_session.flush()

    course = Course(
        curriculum_id=curriculum.id, course_code="EXP101", name_th="วิชาทดสอบส่งออก", credit=3
    )
    plo1 = PLO(curriculum_id=curriculum.id, code="PLO1", description_th="PLO หนึ่ง", category="ความรู้")
    plo2 = PLO(curriculum_id=curriculum.id, code="PLO2", description_th="PLO สอง", category="ทักษะ")
    owner = User(
        username="exp-owner", password="unused", first_name="สมชาย", last_name="ใจดี", role="instructor"
    )
    other = User(
        username="exp-other", password="unused", first_name="อื่น", last_name="ไม่เกี่ยว", role="instructor"
    )
    db_session.add_all([course, plo1, plo2, owner, other])
    db_session.flush()

    offering = CourseOffering(
        course_id=course.id, instructor_id=owner.id, academic_year=2569, semester=1, section="1"
    )
    db_session.add(offering)
    db_session.flush()

    for sid in ("EXP001", "EXP002"):
        db_session.add(
            Student(id=sid, curriculum_id=curriculum.id, first_name="ทดสอบ", last_name=sid, cohort_year=69)
        )
    db_session.flush()
    for sid in ("EXP001", "EXP002"):
        db_session.add(Enrollment(student_id=sid, offering_id=offering.id, final_grade="A"))

    scores = {"CLO-A": {"EXP001": 90, "EXP002": 40}, "CLO-B": {"EXP001": 90}}
    for clo_code, score_by_student in scores.items():
        clo = CLO(
            course_id=course.id,
            code=clo_code,
            description=f"คำอธิบาย {clo_code}",
            pass_threshold_percent=60.00,
            created_by=admin_user_id,
        )
        db_session.add(clo)
        db_session.flush()
        db_session.add(CLOPLOMapping(clo_id=clo.id, plo_id=plo1.id, weight_percent=50.00))
        item = AssessmentItem(offering_id=offering.id, name=f"item-{clo_code}", type="quiz", total_score=100)
        db_session.add(item)
        db_session.flush()
        db_session.add(ItemCLO(item_id=item.id, clo_id=clo.id, weight_percent=100.00))
        for sid, score in score_by_student.items():
            db_session.add(StudentScore(item_id=item.id, student_id=sid, score_obtained=score))

    db_session.commit()
    return {"course": course, "offering": offering, "owner": owner, "other": other}


def test_course_export_has_course_info_and_clo_results(client, db_session, admin_user):
    fx = _make_fixtures(db_session, admin_user.id)

    resp = client.get(f"/export/course/{fx['course'].id}")
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert 'filename="course_EXP101.xlsx"' in resp.headers["content-disposition"]

    wb = load_workbook(io.BytesIO(resp.content))
    assert wb.sheetnames == ["ข้อมูลรายวิชา", "ผล CLO นักศึกษา"]

    info = _sheet_values(wb["ข้อมูลรายวิชา"])
    assert _find_row(info, "รหัสวิชา")[1] == "EXP101"
    assert _find_row(info, "ชื่อวิชา (ไทย)")[1] == "วิชาทดสอบส่งออก"
    assert _find_row(info, "อาจารย์ผู้สอน")[1] == "สมชาย ใจดี"
    assert _find_row(info, "CLO-A")[3] == "PLO1"
    # PLO ของรายวิชามีแค่ PLO1 (PLO2 ไม่มี CLO ของวิชานี้ผูกอยู่)
    assert _find_row(info, "PLO1")[3] == "CLO-A, CLO-B"
    assert not any(row and row[0] == "PLO2" for row in info)
    # ไม่มี Sec ในไฟล์ (ผู้ใช้ขอให้เอาออก)
    assert not any("Sec" in str(cell) for cell in _all_cells(info))

    results = _sheet_values(wb["ผล CLO นักศึกษา"])
    header = _find_row_containing(results, "รหัสนักศึกษา")
    clo_a, clo_b, summary = header.index("CLO-A"), header.index("CLO-B"), header.index("ผ่าน CLO")
    row1 = _find_row_containing(results, "EXP001")
    row2 = _find_row_containing(results, "EXP002")
    assert (row1[clo_a], row1[clo_b], row1[summary]) == ("ผ่าน", "ผ่าน", "2/2")
    # CLO ที่ยังไม่มีคะแนนแยกนับ ไม่ถูกนับรวมเป็น "ไม่ผ่าน"
    assert (row2[clo_a], row2[clo_b], row2[summary]) == ("ไม่ผ่าน", "ไม่มีข้อมูล", "0/2 (ไม่มีข้อมูล 1)")
    assert row1[header.index("ภาคเรียนที่เรียน")] == "2569/1"
    # ไม่มีคะแนนชิ้นงานในไฟล์ (ผู้ใช้ขอแค่ผล ผ่าน/ไม่ผ่าน)
    assert "item-CLO-A" not in _all_cells(results)


def test_course_export_retake_is_one_row_per_enrollment(client, db_session, admin_user):
    """ลงเรียนวิชาเดิมซ้ำคนละภาคเรียน -> แยกแถวตามครั้งที่เรียน ผล CLO คิดจากคะแนนของครั้งนั้นๆ"""
    fx = _make_fixtures(db_session, admin_user.id)
    retake = CourseOffering(
        course_id=fx["course"].id, instructor_id=fx["owner"].id, academic_year=2570, semester=1, section="1"
    )
    db_session.add(retake)
    db_session.flush()
    db_session.add(Enrollment(student_id="EXP002", offering_id=retake.id))
    db_session.commit()

    wb = load_workbook(io.BytesIO(client.get(f"/export/course/{fx['course'].id}").content))
    results = _sheet_values(wb["ผล CLO นักศึกษา"])
    exp002_rows = [row for row in results if "EXP002" in row]
    assert [row[4] for row in exp002_rows] == ["2569/1", "2570/1"]
    # ครั้งที่ 2 ยังไม่มีคะแนนเลย -> ไม่มีข้อมูลทุก CLO (ไม่ปนคะแนนของครั้งแรก)
    assert exp002_rows[1][5:7] == ["ไม่มีข้อมูล", "ไม่มีข้อมูล"]


def test_course_export_labels_group_only_when_term_has_multiple_sections(client, db_session, admin_user):
    """วิชาเปิด 2 กลุ่มในภาคเดียวกัน -> ต่อท้าย "กลุ่ม N" ให้แยกออก ส่วนภาคที่มีกลุ่มเดียวไม่แสดงกลุ่ม"""
    fx = _make_fixtures(db_session, admin_user.id)
    db_session.add_all(
        [
            CourseOffering(
                course_id=fx["course"].id, instructor_id=fx["owner"].id, academic_year=2569, semester=1, section="2"
            ),
            CourseOffering(
                course_id=fx["course"].id, instructor_id=fx["owner"].id, academic_year=2570, semester=1, section="1"
            ),
        ]
    )
    db_session.commit()

    wb = load_workbook(io.BytesIO(client.get(f"/export/course/{fx['course'].id}").content))
    info_cells = _all_cells(_sheet_values(wb["ข้อมูลรายวิชา"]))
    assert {"2569/1 กลุ่ม 1", "2569/1 กลุ่ม 2", "2570/1"} <= set(info_cells)
    assert "2570/1 กลุ่ม 1" not in info_cells
    results = _sheet_values(wb["ผล CLO นักศึกษา"])
    assert _find_row_containing(results, "EXP001")[4] == "2569/1 กลุ่ม 1"


def test_clo_summary_for_course_without_clos_or_scores(client, db_session, admin_user):
    """วิชาที่ยังไม่มี CLO -> "-" (ไม่ใช่ 0/0) และวิชาที่ยังไม่มีคะแนนเลย -> "ไม่มีข้อมูล" (ไม่ใช่ 0/N)"""
    fx = _make_fixtures(db_session, admin_user.id)
    no_clo_course = Course(curriculum_id=fx["course"].curriculum_id, course_code="EXP102", name_th="ยังไม่มี CLO", credit=3)
    db_session.add(no_clo_course)
    db_session.flush()
    no_clo_offering = CourseOffering(
        course_id=no_clo_course.id, instructor_id=fx["owner"].id, academic_year=2569, semester=1, section="1"
    )
    no_score_offering = CourseOffering(
        course_id=fx["course"].id, instructor_id=fx["owner"].id, academic_year=2570, semester=1, section="1"
    )
    db_session.add_all([no_clo_offering, no_score_offering])
    db_session.flush()
    db_session.add_all(
        [
            Enrollment(student_id="EXP001", offering_id=no_clo_offering.id),
            Enrollment(student_id="EXP001", offering_id=no_score_offering.id),
        ]
    )
    db_session.commit()

    courses = _sheet_values(load_workbook(io.BytesIO(client.get("/export/student/EXP001").content))["รายวิชาตามชั้นปี"])
    assert _find_row_containing(courses, "EXP102")[5:7] == ["-", "-"]
    assert _find_row_containing(courses, "2570/1")[5:7] == ["ไม่มีข้อมูล", "ไม่มีข้อมูล"]

    course_wb = load_workbook(io.BytesIO(client.get(f"/export/course/{no_clo_course.id}").content))
    assert _find_row_containing(_sheet_values(course_wb["ผล CLO นักศึกษา"]), "EXP001")[-1] == "-"


def test_course_export_owner_instructor_sees_clo_results(make_client, db_session, admin_user):
    fx = _make_fixtures(db_session, admin_user.id)
    resp = make_client(fx["owner"]).get(f"/export/course/{fx['course'].id}")
    assert resp.status_code == 200
    results = _sheet_values(load_workbook(io.BytesIO(resp.content))["ผล CLO นักศึกษา"])
    assert "EXP001" in _all_cells(results)


def test_course_export_non_owner_instructor_sees_all_student_results(make_client, db_session, admin_user):
    """อาจารย์ที่ไม่ได้สอนวิชานี้ก็เห็นผล CLO ของนักศึกษาทุกกลุ่มเรียน (ผู้ใช้ต้องการให้ดูรายละเอียดได้ครบ)"""
    fx = _make_fixtures(db_session, admin_user.id)
    resp = make_client(fx["other"]).get(f"/export/course/{fx['course'].id}")
    assert resp.status_code == 200
    results = _sheet_values(load_workbook(io.BytesIO(resp.content))["ผล CLO นักศึกษา"])
    header = _find_row_containing(results, "รหัสนักศึกษา")
    assert _find_row_containing(results, "EXP001")[header.index("ผ่าน CLO")] == "2/2"
    assert "EXP002" in _all_cells(results)


def test_student_export_has_profile_plo_and_courses(client, db_session, admin_user):
    _make_fixtures(db_session, admin_user.id)

    resp = client.get("/export/student/EXP001")
    assert resp.status_code == 200
    assert 'filename="student_EXP001.xlsx"' in resp.headers["content-disposition"]

    wb = load_workbook(io.BytesIO(resp.content))
    assert wb.sheetnames == ["ข้อมูลนักศึกษา", "รายวิชาตามชั้นปี", "ผล CLO รายวิชา"]

    info = _sheet_values(wb["ข้อมูลนักศึกษา"])
    assert _find_row(info, "รหัสนักศึกษา")[1] == "EXP001"
    assert _find_row(info, "ชื่อ-นามสกุล")[1] == "ทดสอบ EXP001"
    assert _find_row(info, "รุ่น")[1] == 69
    assert _find_row(info, "หลักสูตร")[1] == "หลักสูตรทดสอบ Export (2569)"
    plo1 = _find_row(info, "PLO1")
    assert (plo1[3], plo1[4]) == (90.0, "บรรลุ")
    # PLO2 ไม่มี CLO ผูกอยู่เลย -> ยังไม่มีข้อมูล (แยกจาก "ยังไม่บรรลุ" เหมือนหน้าเว็บ)
    assert _find_row(info, "PLO2")[4] == "ยังไม่มีข้อมูล"
    assert _find_row(info, "บรรลุ PLO")[1] == "1/2 ข้อ"

    courses = _sheet_values(wb["รายวิชาตามชั้นปี"])
    header = _find_row_containing(courses, "รหัสวิชา")
    assert "Sec" not in header and "เกรด" not in header
    # ไม่มีแผนการเรียน -> วิชาที่ลงทะเบียนไปอยู่ในส่วน "นอกแผน" ไม่หายไปจากไฟล์
    assert "วิชาที่ลงทะเบียนนอกแผนการเรียน" in _all_cells(courses)
    course_row = _find_row_containing(courses, "EXP101")
    assert course_row == ["EXP101", "วิชาทดสอบส่งออก", 3, "2569/1", "สมชาย ใจดี", "ผ่าน", "2/2"]

    clo_rows = _sheet_values(wb["ผล CLO รายวิชา"])
    clo_a = _find_row_containing(clo_rows, "CLO-A")
    assert clo_a == ["2569/1", "EXP101", "วิชาทดสอบส่งออก", "CLO-A", "คำอธิบาย CLO-A", 90.0, 60.0, "ผ่าน"]


def test_student_export_not_achieved(client, db_session, admin_user):
    _make_fixtures(db_session, admin_user.id)

    wb = load_workbook(io.BytesIO(client.get("/export/student/EXP002").content))
    plo1 = _find_row(_sheet_values(wb["ข้อมูลนักศึกษา"]), "PLO1")
    assert (plo1[3], plo1[4]) == (40.0, "ยังไม่บรรลุ")
    course_row = _find_row_containing(_sheet_values(wb["รายวิชาตามชั้นปี"]), "EXP101")
    assert course_row[5:7] == ["ไม่ผ่าน", "0/2 (ไม่มีข้อมูล 1)"]
    clo_rows = _sheet_values(wb["ผล CLO รายวิชา"])
    assert _find_row_containing(clo_rows, "CLO-A")[5:8] == [40.0, 60.0, "ไม่ผ่าน"]
    assert _find_row_containing(clo_rows, "CLO-B")[5:8] == ["-", 60.0, "ไม่มีข้อมูล"]


def test_student_export_groups_courses_by_study_plan_year(client, db_session, admin_user):
    """วิชาในแผนการเรียนแยกตามชั้นปี พร้อมผล YLO — วิชาในแผนที่ยังไม่ลงทะเบียนแสดง "ยังไม่ลงทะเบียน"
    และชั้นปีที่นักศึกษายังไปไม่ถึงแสดง "ยังไม่ถึงชั้นปีนี้" (โครงเดียวกับหน้า /student-plo)"""
    fx = _make_fixtures(db_session, admin_user.id)
    curriculum_id = fx["course"].curriculum_id
    later_course = Course(curriculum_id=curriculum_id, course_code="EXP201", name_th="วิชาปีสอง", credit=2)
    db_session.add(later_course)
    db_session.flush()
    db_session.add_all(
        [
            StudyPlan(curriculum_id=curriculum_id, course_id=fx["course"].id, year_level=1, semester=1),
            StudyPlan(curriculum_id=curriculum_id, course_id=later_course.id, year_level=2, semester=1),
            YLO(curriculum_id=curriculum_id, year_level=1, description="YLO ปีหนึ่ง"),
        ]
    )
    db_session.commit()

    rows = _sheet_values(load_workbook(io.BytesIO(client.get("/export/student/EXP001").content))["รายวิชาตามชั้นปี"])
    year_rows = [row for row in rows if row[0] and str(row[0]).startswith("ชั้นปีที่")]
    assert [row[0] for row in year_rows] == ["ชั้นปีที่ 1", "ชั้นปีที่ 2", "ชั้นปีที่ 3", "ชั้นปีที่ 4"]
    assert _find_row(rows, "YLO")[1] == "YLO ปีหนึ่ง"
    assert _find_row_containing(rows, "EXP101")[5] == "ผ่าน"
    assert _find_row_containing(rows, "EXP201")[3:6] == ["-", "-", "ยังไม่ลงทะเบียน"]
    assert "วิชาที่ลงทะเบียนนอกแผนการเรียน" not in _all_cells(rows)
    # ชั้นปีของรุ่น 69 ขึ้นกับวันที่รันเทส - เช็คเฉพาะปี 4 ที่ยังไปไม่ถึงแน่นอนอีกหลายปี
    assert year_rows[3][1] == "ยังไม่ถึงชั้นปีนี้"


def test_student_export_non_owner_instructor_sees_plo_and_clo(make_client, db_session, admin_user):
    """อาจารย์ที่ไม่ได้สอนวิชานี้ก็เห็นทั้งผล PLO และผล CLO ทุกวิชาของนักศึกษา"""
    fx = _make_fixtures(db_session, admin_user.id)
    resp = make_client(fx["other"]).get("/export/student/EXP001")
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.content))
    assert _find_row(_sheet_values(wb["ข้อมูลนักศึกษา"]), "PLO1")[4] == "บรรลุ"
    assert _find_row_containing(_sheet_values(wb["รายวิชาตามชั้นปี"]), "EXP101")[5:7] == ["ผ่าน", "2/2"]
    assert _find_row_containing(_sheet_values(wb["ผล CLO รายวิชา"]), "CLO-A")[7] == "ผ่าน"


def test_export_not_found_returns_404(client, db_session):
    assert client.get("/export/course/999999").status_code == 404
    assert client.get("/export/student/NOPE999").status_code == 404
