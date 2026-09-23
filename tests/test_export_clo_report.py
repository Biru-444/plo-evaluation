"""
Tests สำหรับ GET /clo-achievement/export - export รายงานผลบรรลุ CLO เป็น Excel (6 ชีต) ของ
course_offering เดียว ครอบคลุมทั้งความถูกต้องของตัวเลข (ต้องตรงกับ GET /clo-achievement ที่ใช้สูตร
เดียวกันจาก clo_achievement_service.py), สิทธิ์การเข้าถึง (admin/อาจารย์เจ้าของ/role อื่น), และ
edge case (CLO ไม่มีชิ้นงานผูก, ยังไม่มีเกรด)

ใช้ fixture pattern เดียวกับ tests/test_courses_enrolled_students_mastery.py (สร้างข้อมูลเองใน
db_session ทุกเทส ไม่พึ่งข้อมูลที่มีอยู่ก่อน - rollback อัตโนมัติหลังจบเทสตาม conftest.py)
"""
from __future__ import annotations

import re
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.auth import get_current_user
from app.database import get_db
from app.main import app
from app.models import (
    CLO,
    AssessmentItem,
    Course,
    CourseOffering,
    Curriculum,
    Enrollment,
    ItemCLO,
    Student,
    StudentScore,
    User,
)


@pytest.fixture()
def make_client(db_session):
    """เหมือน conftest.py's `client` fixture แต่รับ user ที่จะให้เป็น current_user ได้ (ไม่ hardcode
    admin_user) - ใช้ทดสอบสิทธิ์ที่ต้องสวมบทเป็น user คนอื่นที่ไม่ใช่ admin"""

    def _make(user: User) -> TestClient:
        def _override_get_db():
            yield db_session

        def _override_get_current_user():
            return user

        app.dependency_overrides[get_db] = _override_get_db
        app.dependency_overrides[get_current_user] = _override_get_current_user
        return TestClient(app)

    yield _make
    app.dependency_overrides.clear()


def _make_curriculum_course_offering(db_session, instructor=None, section="1"):
    curriculum = Curriculum(name="Test Curriculum CLO Report", year=2569)
    db_session.add(curriculum)
    db_session.flush()

    course = Course(
        curriculum_id=curriculum.id,
        course_code="CLORPT1",
        name_th="วิชาทดสอบรายงาน CLO",
        name_en="CLO Report Test Course",
        credit=3,
        category="วิชาแกน",
    )
    db_session.add(course)
    db_session.flush()

    offering = CourseOffering(
        course_id=course.id,
        instructor_id=instructor.id if instructor is not None else None,
        academic_year=2569,
        semester=1,
        section=section,
    )
    db_session.add(offering)
    db_session.flush()

    return curriculum, course, offering


def _enroll_student(db_session, offering, student_id, final_grade=None):
    student = Student(
        id=student_id,
        curriculum_id=offering.course.curriculum_id,
        first_name="ทดสอบ",
        last_name=student_id,
        cohort_year=69,
        current_year_level=1,
    )
    db_session.add(student)
    db_session.add(Enrollment(student_id=student.id, offering_id=offering.id, final_grade=final_grade))
    db_session.flush()
    return student


def _add_clo_with_score(db_session, *, course, offering, clo_code, admin_user_id, scores: dict):
    """สร้าง CLO 1 ตัว + assessment_item 1 ชิ้น (เต็ม 100, weight 100%) - scores คือ
    {student_id: score_obtained} ใครไม่อยู่ใน dict = ไม่มีคะแนนเลย (จำลอง "ไม่มีข้อมูล")"""
    clo = CLO(
        course_id=course.id,
        code=clo_code,
        description=f"ทดสอบ {clo_code}",
        pass_threshold_percent=60.00,
        created_by=admin_user_id,
    )
    db_session.add(clo)
    db_session.flush()

    item = AssessmentItem(offering_id=offering.id, name=f"item-{clo_code}", type="quiz", total_score=100.0)
    db_session.add(item)
    db_session.flush()
    db_session.add(ItemCLO(item_id=item.id, clo_id=clo.id, weight_percent=100.00))

    for student_id, score in scores.items():
        db_session.add(StudentScore(item_id=item.id, student_id=student_id, score_obtained=score))

    return clo


def _sheet_clo_rows(wb):
    ws = wb["ผลการบรรลุ CLO"]
    rows = list(ws.iter_rows(min_row=2, values_only=False))
    return {row[0].value: row for row in rows}  # keyed by CLO code


ALL_SHEET_NAMES = [
    "คำอธิบาย",
    "ข้อมูลรายวิชา",
    "ผลการบรรลุ CLO",
    "สรุปผลการเรียน",
    "การยืนยันผลสัมฤทธิ์",
    "รายบุคคล",
]


class TestExportSucceedsWithSixSheets:
    def test_export_returns_xlsx_with_six_sheets(self, client, db_session, admin_user):
        curriculum, course, offering = _make_curriculum_course_offering(db_session)
        student = _enroll_student(db_session, offering, "CLORPT-001", final_grade="A")
        _add_clo_with_score(
            db_session,
            course=course,
            offering=offering,
            clo_code="CLO1",
            admin_user_id=admin_user.id,
            scores={student.id: 90.0},
        )
        db_session.commit()

        resp = client.get(f"/clo-achievement/export?offering_id={offering.id}")
        assert resp.status_code == 200
        assert resp.headers["content-type"] == (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        assert "clo_report_CLORPT1" in resp.headers["content-disposition"]

        wb = load_workbook(BytesIO(resp.content))
        assert wb.sheetnames == ALL_SHEET_NAMES


class TestExplanationSheet:
    def test_explanation_sheet_mentions_target_rate_and_has_timestamp(
        self, client, db_session, admin_user
    ):
        curriculum, course, offering = _make_curriculum_course_offering(db_session)
        db_session.commit()

        resp = client.get(f"/clo-achievement/export?offering_id={offering.id}&target_rate=75")
        assert resp.status_code == 200
        wb = load_workbook(BytesIO(resp.content))
        ws = wb["คำอธิบาย"]

        labels = [ws.cell(row=r, column=1).value for r in range(3, ws.max_row + 1)]
        values = [ws.cell(row=r, column=2).value for r in range(3, ws.max_row + 1)]
        rows_by_label = dict(zip(labels, values))

        assert any("เกณฑ์ระดับรายวิชา" in label for label in labels)
        target_rate_value = next(v for k, v in rows_by_label.items() if "เกณฑ์ระดับรายวิชา" in k)
        assert "75" in target_rate_value

        assert any("export" in label.lower() for label in labels)
        assert any("ผ่าน" in label for label in labels)
        assert any("บรรลุ" in label for label in labels)
        assert any("สูตร" in label for label in labels)

        # ไม่ยึดตามแบบฟอร์มราชการ - ต้องไม่มีคำอ้างอิงหมวด/ข้อแบบ มคอ.5 (เช่น "หมวด 1", "หมวด 3 ข้อ 1-4")
        # หลงเหลืออยู่ในชีตไหนเลย ("หมวดวิชา" ซึ่งเป็นชื่อฟิลด์ category ของวิชาเองไม่นับ - เป็นคำละ
        # ความหมายกันคนละเรื่อง)
        for name in wb.sheetnames:
            ws2 = wb[name]
            for row in ws2.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str):
                        assert "มคอ." not in cell.value
                        assert "OBE5" not in cell.value
                        assert not re.search(r"หมวด\s*\d", cell.value)


class TestSheetMatchesGetCloAchievement:
    def test_clo_sheet_numbers_match_endpoint(self, client, db_session, admin_user):
        curriculum, course, offering = _make_curriculum_course_offering(db_session)
        s1 = _enroll_student(db_session, offering, "CLORPT-010")
        s2 = _enroll_student(db_session, offering, "CLORPT-011")
        s3 = _enroll_student(db_session, offering, "CLORPT-012")
        _add_clo_with_score(
            db_session,
            course=course,
            offering=offering,
            clo_code="CLO1",
            admin_user_id=admin_user.id,
            scores={s1.id: 90.0, s2.id: 30.0},  # s3 ไม่มีคะแนนเลย
        )
        db_session.commit()

        api_resp = client.get(f"/clo-achievement?offering_id={offering.id}")
        assert api_resp.status_code == 200
        api_clo = api_resp.json()["clo_achievements"][0]

        export_resp = client.get(f"/clo-achievement/export?offering_id={offering.id}")
        wb = load_workbook(BytesIO(export_resp.content))
        rows = _sheet_clo_rows(wb)
        row = rows["CLO1"]

        assert row[6].value == api_clo["passed_count"]  # col 7: ผ่าน (คน)
        assert row[7].value == api_clo["failed_count"]  # col 8: ไม่ผ่าน (คน)
        assert row[8].value == api_clo["students_without_data"]  # col 9: ไม่มีข้อมูล (คน)
        assert row[9].value == api_clo["achieved_rate_percent"]  # col 10: ร้อยละที่ผ่าน


class TestStatusChangesWithTargetRate:
    def test_status_flips_between_low_and_high_target_rate(self, client, db_session, admin_user):
        curriculum, course, offering = _make_curriculum_course_offering(db_session)
        s1 = _enroll_student(db_session, offering, "CLORPT-020")
        s2 = _enroll_student(db_session, offering, "CLORPT-021")
        # 1 ผ่าน 1 ไม่ผ่าน -> achieved_rate_percent = 50.0
        _add_clo_with_score(
            db_session,
            course=course,
            offering=offering,
            clo_code="CLO1",
            admin_user_id=admin_user.id,
            scores={s1.id: 90.0, s2.id: 10.0},
        )
        db_session.commit()

        low_resp = client.get(f"/clo-achievement/export?offering_id={offering.id}&target_rate=40")
        high_resp = client.get(f"/clo-achievement/export?offering_id={offering.id}&target_rate=90")

        low_status = _sheet_clo_rows(load_workbook(BytesIO(low_resp.content)))["CLO1"][10].value
        high_status = _sheet_clo_rows(load_workbook(BytesIO(high_resp.content)))["CLO1"][10].value

        assert low_status == "บรรลุ"
        assert high_status == "ไม่บรรลุ"


class TestGradeDistribution:
    def test_percentage_and_total_row_are_correct(self, client, db_session, admin_user):
        curriculum, course, offering = _make_curriculum_course_offering(db_session)
        # 45 คนทั้งหมด, 3 คนได้ I -> 3/45 = 6.666...% ต้องปัดเป็น 6.67
        for i in range(42):
            _enroll_student(db_session, offering, f"CLORPT-G{i:03d}", final_grade="A")
        for i in range(3):
            _enroll_student(db_session, offering, f"CLORPT-GI{i:03d}", final_grade="I")
        db_session.commit()

        resp = client.get(f"/clo-achievement/export?offering_id={offering.id}")
        wb = load_workbook(BytesIO(resp.content))
        ws = wb["สรุปผลการเรียน"]

        grade_rows = {row[0].value: row for row in ws.iter_rows(min_row=9, values_only=False)}
        assert grade_rows["I"][2].value == 3
        assert grade_rows["I"][3].value == "6.67"

        total_row = grade_rows["รวม"]
        assert total_row[2].value == 45
        assert total_row[3].value == "100.00"

    def test_withdrawn_excluded_from_remaining(self, client, db_session, admin_user):
        curriculum, course, offering = _make_curriculum_course_offering(db_session)
        for i in range(8):
            _enroll_student(db_session, offering, f"CLORPT-R{i:03d}", final_grade="B")
        for i in range(2):
            _enroll_student(db_session, offering, f"CLORPT-W{i:03d}", final_grade="W")
        db_session.commit()

        resp = client.get(f"/clo-achievement/export?offering_id={offering.id}")
        wb = load_workbook(BytesIO(resp.content))
        ws = wb["สรุปผลการเรียน"]

        assert ws.cell(row=3, column=2).value == 10  # จำนวนลงทะเบียน
        assert ws.cell(row=4, column=2).value == 2  # จำนวนที่ถอน (W)
        assert ws.cell(row=5, column=2).value == 8  # คงอยู่เมื่อสิ้นภาค = 10 - 2

    def test_no_grades_at_all_shows_note(self, client, db_session, admin_user):
        curriculum, course, offering = _make_curriculum_course_offering(db_session)
        _enroll_student(db_session, offering, "CLORPT-NG1", final_grade=None)
        db_session.commit()

        resp = client.get(f"/clo-achievement/export?offering_id={offering.id}")
        wb = load_workbook(BytesIO(resp.content))
        ws = wb["สรุปผลการเรียน"]
        assert "ยังไม่มีข้อมูลเกรด" in ws.cell(row=6, column=1).value


class TestAccessControl:
    def test_instructor_not_owner_returns_403(self, make_client, db_session, admin_user):
        owner = User(
            username="clorpt-owner", password="x", first_name="เจ้าของ", last_name="วิชา", role="instructor"
        )
        other = User(
            username="clorpt-other", password="x", first_name="คนอื่น", last_name="ไม่เกี่ยว", role="instructor"
        )
        db_session.add_all([owner, other])
        db_session.flush()
        curriculum, course, offering = _make_curriculum_course_offering(db_session, instructor=owner)
        db_session.commit()

        resp = make_client(other).get(f"/clo-achievement/export?offering_id={offering.id}")
        assert resp.status_code == 403

    def test_owning_instructor_can_export(self, make_client, db_session, admin_user):
        owner = User(
            username="clorpt-owner2", password="x", first_name="เจ้าของ", last_name="วิชา", role="instructor"
        )
        db_session.add(owner)
        db_session.flush()
        curriculum, course, offering = _make_curriculum_course_offering(db_session, instructor=owner)
        db_session.commit()

        resp = make_client(owner).get(f"/clo-achievement/export?offering_id={offering.id}")
        assert resp.status_code == 200

    def test_nonexistent_offering_returns_404(self, client):
        resp = client.get("/clo-achievement/export?offering_id=999999")
        assert resp.status_code == 404

    def test_other_role_gets_no_personal_sheet(self, make_client, db_session, admin_user):
        """role ที่ไม่ใช่ admin/instructor (เช่น ประธานหลักสูตร ถ้ามีในอนาคต) - ได้ไฟล์แต่ไม่มีชีต
        รายบุคคล (ข้อมูล PDPA)"""
        coordinator = User(
            username="clorpt-coordinator",
            password="x",
            first_name="ประธาน",
            last_name="หลักสูตร",
            role="coordinator",
        )
        db_session.add(coordinator)
        db_session.flush()
        curriculum, course, offering = _make_curriculum_course_offering(db_session)
        db_session.commit()

        resp = make_client(coordinator).get(f"/clo-achievement/export?offering_id={offering.id}")
        assert resp.status_code == 200
        wb = load_workbook(BytesIO(resp.content))
        assert "รายบุคคล" not in wb.sheetnames
        assert len(wb.sheetnames) == 5

    def test_admin_gets_personal_sheet(self, client, db_session):
        curriculum, course, offering = _make_curriculum_course_offering(db_session)
        db_session.commit()

        resp = client.get(f"/clo-achievement/export?offering_id={offering.id}")
        assert resp.status_code == 200
        wb = load_workbook(BytesIO(resp.content))
        assert "รายบุคคล" in wb.sheetnames


class TestCloWithoutAssessmentItem:
    def test_clo_without_any_item_shows_no_data_status_not_error(self, db_session, admin_user, client):
        curriculum, course, offering = _make_curriculum_course_offering(db_session)
        _enroll_student(db_session, offering, "CLORPT-030")
        clo = CLO(
            course_id=course.id,
            code="CLO-NOITEM",
            description="CLO ไม่มีชิ้นงานผูก",
            pass_threshold_percent=60.00,
            created_by=admin_user.id,
        )
        db_session.add(clo)
        db_session.commit()

        resp = client.get(f"/clo-achievement/export?offering_id={offering.id}")
        assert resp.status_code == 200
        wb = load_workbook(BytesIO(resp.content))
        row = _sheet_clo_rows(wb)["CLO-NOITEM"]
        assert row[10].value == "ไม่มีข้อมูล"  # col 11: สถานะ
