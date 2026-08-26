"""API routes for Enrollment"""
from __future__ import annotations

import csv
import io

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from openpyxl import load_workbook
from sqlalchemy.dialects.postgresql import insert as pg_insert
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.auth import get_current_user
from app.database import get_db
from app.models import Course, CourseOffering, Enrollment, Student, User
from app.schemas import (
    BulkEnrollByCohortResult,
    BulkEnrollByCohortSchema,
    BulkEnrollResult,
    BulkEnrollSchema,
    EnrolledStudentBrief,
    EnrollmentCreateSchema,
    EnrollmentSchema,
    EnrollmentUpdateSchema,
    OtherSectionConflict,
)

router = APIRouter(prefix="/enrollments", tags=["Enrollments"])

STUDENT_ID_HEADER_ALIASES = {"รหัสนักศึกษา", "student_id", "id"}


def _require_offering_ownership(
    db: Session, offering_id: int, current_user: User
) -> CourseOffering:
    offering = db.get(CourseOffering, offering_id)
    if offering is None:
        raise HTTPException(status_code=404, detail="Course offering not found")
    if current_user.role != "admin" and offering.instructor_id != current_user.id:
        raise HTTPException(status_code=403, detail="คุณไม่ใช่ผู้สอนวิชานี้")
    return offering


def _sibling_offering_ids(db: Session, offering: CourseOffering) -> list[int]:
    """Other CourseOffering rows for the same course/term (different section/หมู่)."""
    rows = (
        db.query(CourseOffering.id)
        .filter(
            CourseOffering.course_id == offering.course_id,
            CourseOffering.academic_year == offering.academic_year,
            CourseOffering.semester == offering.semester,
            CourseOffering.id != offering.id,
        )
        .all()
    )
    return [row[0] for row in rows]


def _students_in_other_sections(
    db: Session, offering: CourseOffering, student_ids: list[str]
) -> dict[str, str]:
    """Of the given student_ids, which are already enrolled in a *sibling* section
    (same course/academic_year/semester, different offering) - mapped to that
    section's label. Cohort 69 style courses split into multiple sections/หมู่ for
    room capacity, so bulk-add-by-cohort must not re-lump a student already correctly
    placed in the other section into this one."""
    sibling_ids = _sibling_offering_ids(db, offering)
    if not sibling_ids or not student_ids:
        return {}
    rows = (
        db.query(Enrollment.student_id, CourseOffering.section)
        .join(CourseOffering, CourseOffering.id == Enrollment.offering_id)
        .filter(Enrollment.offering_id.in_(sibling_ids), Enrollment.student_id.in_(student_ids))
        .all()
    )
    return {student_id: section for student_id, section in rows}


def _bulk_enroll(db: Session, offering: CourseOffering, raw_student_ids: list[str]) -> BulkEnrollResult:
    """Shared insert logic for /bulk and /bulk-upload - both report the same
    per-student outcome breakdown (added / already enrolled / not found /
    wrong curriculum), just sourced from a JSON list vs a parsed file."""
    seen: set[str] = set()
    requested_ids: list[str] = []
    for sid in raw_student_ids:
        sid = (sid or "").strip()
        if sid and sid not in seen:
            seen.add(sid)
            requested_ids.append(sid)

    course = db.get(Course, offering.course_id)

    students = db.query(Student).filter(Student.id.in_(requested_ids)).all()
    student_by_id = {s.id: s for s in students}

    not_found = [sid for sid in requested_ids if sid not in student_by_id]
    wrong_curriculum = [
        sid
        for sid in requested_ids
        if sid in student_by_id and student_by_id[sid].curriculum_id != course.curriculum_id
    ]
    valid_ids = [
        sid
        for sid in requested_ids
        if sid in student_by_id and student_by_id[sid].curriculum_id == course.curriculum_id
    ]

    already_enrolled_ids: set[str] = set()
    if valid_ids:
        already_enrolled_ids = {
            row[0]
            for row in db.query(Enrollment.student_id)
            .filter(Enrollment.offering_id == offering.id, Enrollment.student_id.in_(valid_ids))
            .all()
        }
    not_yet_here = [sid for sid in valid_ids if sid not in already_enrolled_ids]
    other_section_map = _students_in_other_sections(db, offering, not_yet_here)
    to_insert = [sid for sid in not_yet_here if sid not in other_section_map]

    added_count = 0
    if to_insert:
        stmt = (
            pg_insert(Enrollment.__table__)
            .values([{"student_id": sid, "offering_id": offering.id} for sid in to_insert])
            .on_conflict_do_nothing(index_elements=["student_id", "offering_id"])
            .returning(Enrollment.__table__.c.student_id)
        )
        result = db.execute(stmt)
        added_count = len(result.fetchall())
        db.commit()

    return BulkEnrollResult(
        added_count=added_count,
        already_enrolled=sorted(already_enrolled_ids),
        not_found=not_found,
        wrong_curriculum=wrong_curriculum,
        already_in_other_section=[
            OtherSectionConflict(student_id=sid, section=section)
            for sid, section in sorted(other_section_map.items())
        ],
    )


def _parse_roster_file(filename: str, content: bytes) -> list[str]:
    lower_name = filename.lower()
    if lower_name.endswith(".csv"):
        rows = list(csv.reader(io.StringIO(content.decode("utf-8-sig"))))
    elif lower_name.endswith(".xlsx"):
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        worksheet = workbook.active
        rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
    else:
        raise HTTPException(status_code=400, detail="รองรับเฉพาะไฟล์ .csv หรือ .xlsx เท่านั้น")

    rows = [row for row in rows if row and any(cell not in (None, "") for cell in row)]
    if not rows:
        raise HTTPException(status_code=400, detail="ไฟล์ว่างเปล่า")

    header = [str(cell).strip().lower() if cell is not None else "" for cell in rows[0]]
    col_index = next(
        (i for i, h in enumerate(header) if h in STUDENT_ID_HEADER_ALIASES), None
    )
    if col_index is None:
        raise HTTPException(
            status_code=400,
            detail="ไม่พบคอลัมน์รหัสนักศึกษา - หัวคอลัมน์ต้องชื่อ 'รหัสนักศึกษา', 'student_id', หรือ 'id'",
        )

    student_ids: list[str] = []
    for row in rows[1:]:
        if col_index >= len(row) or row[col_index] is None:
            continue
        raw = row[col_index]
        if isinstance(raw, float) and raw.is_integer():
            sid = str(int(raw))
        else:
            sid = str(raw).strip()
        if sid:
            student_ids.append(sid)
    return student_ids


@router.get("", response_model=list[EnrollmentSchema])
def list_enrollments(
    offering_id: int | None = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    query = db.query(Enrollment)
    if offering_id is not None:
        query = query.filter(Enrollment.offering_id == offering_id)
    return query.order_by(Enrollment.id).all()


@router.get("/sibling-sections", response_model=list[OtherSectionConflict])
def list_sibling_section_enrollments(
    offering_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """For a course split into multiple หมู่/section (e.g. รุ่น 69 with 2 sections),
    tells the caller which students are already enrolled in a *different* offering of
    the same course/term - used by the enrollment UI to label/exclude them instead of
    silently lumping both sections' students together. Registered before /{enrollment_id}
    so it isn't swallowed by that route (FastAPI matches path routes in registration order,
    and an untyped/str path param would otherwise greedily match "sibling-sections" too)."""
    offering = _require_offering_ownership(db, offering_id, current_user)
    sibling_ids = _sibling_offering_ids(db, offering)
    if not sibling_ids:
        return []
    rows = (
        db.query(Enrollment.student_id, CourseOffering.section)
        .join(CourseOffering, CourseOffering.id == Enrollment.offering_id)
        .filter(Enrollment.offering_id.in_(sibling_ids))
        .all()
    )
    return [OtherSectionConflict(student_id=sid, section=section) for sid, section in rows]


@router.get("/{enrollment_id}", response_model=EnrollmentSchema)
def get_enrollment(
    enrollment_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    enrollment = db.get(Enrollment, enrollment_id)
    if enrollment is None:
        raise HTTPException(status_code=404, detail="Enrollment not found")
    return enrollment


@router.post("", response_model=EnrollmentSchema, status_code=201)
def create_enrollment(
    payload: EnrollmentCreateSchema,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    if current_user.role != "admin":
        offering = db.get(CourseOffering, payload.offering_id)
        if offering is None or offering.instructor_id != current_user.id:
            raise HTTPException(status_code=403, detail="คุณไม่ใช่ผู้สอนวิชานี้")

    enrollment = Enrollment(**payload.model_dump())
    db.add(enrollment)
    try:
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(
            status_code=409,
            detail="Enrollment already exists or references an invalid student/offering",
        ) from exc
    db.refresh(enrollment)
    return enrollment


@router.put("/{enrollment_id}", response_model=EnrollmentSchema)
def update_enrollment(
    enrollment_id: int,
    payload: EnrollmentUpdateSchema,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    enrollment = db.get(Enrollment, enrollment_id)
    if enrollment is None:
        raise HTTPException(status_code=404, detail="Enrollment not found")
    if current_user.role != "admin":
        offering = db.get(CourseOffering, enrollment.offering_id)
        if offering is None or offering.instructor_id != current_user.id:
            raise HTTPException(status_code=403, detail="คุณไม่ใช่ผู้สอนวิชานี้")
    for field, value in payload.model_dump(exclude_unset=True).items():
        setattr(enrollment, field, value)
    try:
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(status_code=409, detail="Enrollment could not be updated") from exc
    db.refresh(enrollment)
    return enrollment


@router.delete("/{enrollment_id}", status_code=204)
def delete_enrollment(
    enrollment_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    enrollment = db.get(Enrollment, enrollment_id)
    if enrollment is None:
        raise HTTPException(status_code=404, detail="Enrollment not found")
    if current_user.role != "admin":
        offering = db.get(CourseOffering, enrollment.offering_id)
        if offering is None or offering.instructor_id != current_user.id:
            raise HTTPException(status_code=403, detail="คุณไม่ใช่ผู้สอนวิชานี้")
    db.delete(enrollment)  # cascade ลบ ไม่มี child table อ้างถึง enrollment โดยตรง
    db.commit()


@router.post("/bulk-by-cohort", response_model=BulkEnrollByCohortResult)
def bulk_enroll_by_cohort(
    payload: BulkEnrollByCohortSchema,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    offering = _require_offering_ownership(db, payload.offering_id, current_user)
    course = db.get(Course, offering.course_id)

    students = (
        db.query(Student)
        .filter(
            Student.cohort_year == payload.cohort_year,
            Student.curriculum_id == course.curriculum_id,
        )
        .order_by(Student.id)
        .all()
    )

    if not students:
        return BulkEnrollByCohortResult(added_count=0, already_enrolled_count=0, added_students=[])

    existing_ids = {
        row[0]
        for row in db.query(Enrollment.student_id)
        .filter(Enrollment.offering_id == payload.offering_id)
        .all()
    }

    not_yet_here = [s for s in students if s.id not in existing_ids]
    other_section_map = _students_in_other_sections(db, offering, [s.id for s in not_yet_here])
    candidates = [s for s in not_yet_here if s.id not in other_section_map]
    already_count = len(students) - len(not_yet_here)

    added_students: list[Student] = []
    if candidates:
        stmt = (
            pg_insert(Enrollment.__table__)
            .values([{"student_id": s.id, "offering_id": payload.offering_id} for s in candidates])
            .on_conflict_do_nothing(index_elements=["student_id", "offering_id"])
            .returning(Enrollment.__table__.c.student_id)
        )
        result = db.execute(stmt)
        added_ids = {row[0] for row in result.fetchall()}
        db.commit()
        added_students = [s for s in candidates if s.id in added_ids]
        already_count += len(candidates) - len(added_students)

    return BulkEnrollByCohortResult(
        added_count=len(added_students),
        already_enrolled_count=already_count,
        added_students=[
            EnrolledStudentBrief(id=s.id, first_name=s.first_name, last_name=s.last_name)
            for s in added_students
        ],
        already_in_other_section=[
            OtherSectionConflict(student_id=sid, section=section)
            for sid, section in sorted(other_section_map.items())
        ],
    )


@router.post("/bulk", response_model=BulkEnrollResult)
def bulk_enroll_students(
    payload: BulkEnrollSchema,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    offering = _require_offering_ownership(db, payload.offering_id, current_user)
    return _bulk_enroll(db, offering, payload.student_ids)


@router.post("/bulk-upload", response_model=BulkEnrollResult)
async def bulk_enroll_upload(
    offering_id: int = Form(...),
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    offering = _require_offering_ownership(db, offering_id, current_user)
    content = await file.read()
    student_ids = _parse_roster_file(file.filename or "", content)
    return _bulk_enroll(db, offering, student_ids)
