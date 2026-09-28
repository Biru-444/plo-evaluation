"""
ทำอะไร : ส่งออกข้อมูลเป็นไฟล์ Excel (.xlsx) 2 แบบ
         - GET /export/course/{course_id} : ไฟล์ของรายวิชา 1 วิชา มี 2 ชีท
             1. "ข้อมูลรายวิชา" — รหัส/ชื่อวิชา, อาจารย์ผู้สอน (ทุกกลุ่มเรียนที่เปิดสอน), CLO ของวิชา และ
                PLO ที่วิชานี้รับผิดชอบ
             2. "ผล CLO นักศึกษา" — รายชื่อนักศึกษาที่ลงทะเบียนวิชานี้ + ผ่าน/ไม่ผ่านรายข้อ CLO
                (ไม่มีคะแนนชิ้นงาน ตามที่ผู้ใช้ขอ) 1 แถวต่อ 1 การลงทะเบียน (ลงเรียนซ้ำ = แยกแถวตามครั้ง)
         - GET /export/student/{student_id} : ไฟล์รายบุคคลของนักศึกษา 1 คน มี 2 ชีท
             1. "ข้อมูลนักศึกษา" — ชื่อ, รหัส, รุ่น, หลักสูตร + ผลการบรรลุ PLO ทุกข้อ
             2. "รายวิชาที่เรียน" — ทุกวิชาที่ลงทะเบียน พร้อมจำนวน CLO ที่ผ่านในแต่ละวิชา

เชื่อมกับ : - ผล CLO ใช้ compute_offering_clo_achievement_raw() (clo_achievement_service.py) ตัวเดียวกับ
              GET /clo-achievement และ export รายงาน CLO — คำนวณต่อกลุ่มเรียน (offering)
            - ผล PLO ใช้ _build_plo_requirements + _calculate_plo_achievement_for_student
              (plo_achievement_service.py) ตัวเดียวกับ GET /plo/achievement ที่หน้า /student-plo ใช้
            - ปุ่มดาวน์โหลดอยู่ที่หน้า /curriculum (CurriculumCourses.jsx) และ /student-plo
              (PLOAchievement.jsx)

ถ้าแก้ : นโยบายสิทธิ์ (เดียวกับ commit "Document and enforce authorization policy for score-returning
         endpoints") — ผล CLO รายนักศึกษาเป็นข้อมูลระดับรายวิชา: admin เห็นทุกกลุ่มเรียน อาจารย์เห็นเฉพาะ
         กลุ่มเรียนที่ตัวเองสอน (กรองแถว ไม่ใช่ 403 เพราะ 1 ไฟล์ครอบหลายกลุ่มเรียน) ส่วนผลบรรลุ PLO เป็น
         ข้อมูลระดับหลักสูตร เปิดให้ admin และอาจารย์ทุกคนโดยตั้งใจ
         PLO ของรายวิชามาจาก clo_plo_mapping (แหล่งเดียวกับที่ใช้คำนวณผลบรรลุ PLO) ไม่ใช่ course_plo
"""
from __future__ import annotations

import io

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy.orm import Session

from app.auth import require_role
from app.database import get_db
from app.models import (
    CLO,
    CLOPLOMapping,
    Course,
    CourseOffering,
    Curriculum,
    Enrollment,
    PLO,
    Student,
    User,
)
from app.services.clo_achievement_service import compute_offering_clo_achievement_raw
from app.services.plo_achievement_service import (
    _build_plo_requirements,
    _calculate_plo_achievement_for_student,
)

router = APIRouter(prefix="/export", tags=["Export"])

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# --- สไตล์ที่ใช้ร่วมกันทุกชีท ---
_TITLE_FONT = Font(bold=True, size=14)
_SECTION_FONT = Font(bold=True, size=12)
_LABEL_FONT = Font(bold=True)
_NOTE_FONT = Font(italic=True, color="7F7F7F")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="5B3F8C")
_PASS_FILL = PatternFill("solid", fgColor="D9F2E3")
_FAIL_FILL = PatternFill("solid", fgColor="FBE0E0")
_NO_DATA_FILL = PatternFill("solid", fgColor="EEEEEE")
_THIN = Side(style="thin", color="BFBFBF")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_WRAP = Alignment(wrap_text=True, vertical="top")
_CENTER = Alignment(horizontal="center", vertical="top")

PASS_TEXT = "ผ่าน"
FAIL_TEXT = "ไม่ผ่าน"
NO_DATA_TEXT = "ไม่มีข้อมูล"
PLO_ACHIEVED_TEXT = "บรรลุ"
PLO_NOT_ACHIEVED_TEXT = "ยังไม่บรรลุ"
PLO_NO_DATA_TEXT = "ยังไม่มีข้อมูล"
_STATUS_FILL = {
    PASS_TEXT: _PASS_FILL,
    FAIL_TEXT: _FAIL_FILL,
    NO_DATA_TEXT: _NO_DATA_FILL,
    PLO_ACHIEVED_TEXT: _PASS_FILL,
    PLO_NOT_ACHIEVED_TEXT: _FAIL_FILL,
    PLO_NO_DATA_TEXT: _NO_DATA_FILL,
}
OWN_OFFERINGS_NOTE = "* แสดงผล CLO เฉพาะกลุ่มเรียนที่คุณเป็นผู้สอน"


def _full_name(first_name: str, last_name: str, title: str | None = None) -> str:
    return f"{title or ''}{first_name} {last_name}".strip()


def _term_label(offering: CourseOffering) -> str:
    return f"{offering.academic_year}/{offering.semester}"


def _can_see_offering_scores(offering: CourseOffering, current_user: User) -> bool:
    """ผล CLO รายนักศึกษาเป็นข้อมูลระดับรายวิชา — admin หรืออาจารย์ผู้สอนกลุ่มเรียนนั้นเท่านั้น"""
    return current_user.role == "admin" or offering.instructor_id == current_user.id


def _write_key_values(ws: Worksheet, start_row: int, rows: list[tuple[str, object]]) -> int:
    """เขียนบล็อก "หัวข้อ : ค่า" ทีละแถว (คอลัมน์ A = หัวข้อ ตัวหนา, B = ค่า) คืนแถวว่างถัดไป"""
    row = start_row
    for label, value in rows:
        ws.cell(row=row, column=1, value=label).font = _LABEL_FONT
        ws.cell(row=row, column=2, value=value).alignment = _WRAP
        row += 1
    return row


def _write_table(
    ws: Worksheet,
    start_row: int,
    headers: list[str],
    rows: list[list[object]],
    status_columns: set[int] | None = None,
) -> int:
    """เขียนตารางพร้อมหัวตาราง + เส้นขอบ — status_columns (index เริ่ม 0) คือคอลัมน์ที่เป็นค่าสถานะ
    (ผ่าน/ไม่ผ่าน/ไม่มีข้อมูล, บรรลุ/ยังไม่บรรลุ) จะถูกระบายสีพื้นหลังตามค่า คืนแถวว่างถัดไปหลังตาราง"""
    status_columns = status_columns or set()
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=col, value=header)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.border = _BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    row = start_row + 1
    for values in rows:
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.border = _BORDER
            if (col - 1) in status_columns:
                cell.alignment = _CENTER
                fill = _STATUS_FILL.get(value)
                if fill is not None:
                    cell.fill = fill
            else:
                cell.alignment = _WRAP
        row += 1
    return row


def _set_column_widths(ws: Worksheet, widths: list[int]) -> None:
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width


def _xlsx_response(workbook: Workbook, filename: str) -> Response:
    """แปลง workbook เป็น bytes แล้วส่งกลับเป็นไฟล์แนบ — ชื่อไฟล์ ASCII รูปแบบเดียวกับ export อื่นในระบบ
    (frontend อ่านชื่อไฟล์จาก Content-Disposition ผ่าน _downloadBlob ใน client.js)"""
    buffer = io.BytesIO()
    workbook.save(buffer)
    return Response(
        content=buffer.getvalue(),
        media_type=XLSX_MEDIA_TYPE,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/course/{course_id}")
def export_course_excel(
    course_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_role("admin", "instructor")),
):
    """
    ทำอะไร : สร้างไฟล์ Excel ของรายวิชา 1 วิชา (2 ชีท — ดู docstring ของไฟล์)

    เชื่อมกับ : ชีทที่ 2 คำนวณผล CLO ทีละกลุ่มเรียนด้วย compute_offering_clo_achievement_raw() (ตัวเดียว
                กับ GET /clo-achievement) — "ไม่มีข้อมูล" = ยังไม่มีคะแนนของ CLO นั้นเลย แยกจาก "ไม่ผ่าน"

    ถ้าแก้ : อาจารย์เห็นผล CLO เฉพาะนักศึกษาในกลุ่มเรียนที่ตัวเองสอน (ดูนโยบายสิทธิ์ใน docstring ของไฟล์)
             ชีทที่ 1 (ข้อมูลรายวิชา) ไม่มีข้อมูลคะแนน จึงแสดงครบทุกกลุ่มเรียน — 404 ถ้าไม่พบรายวิชา
    """
    course = db.get(Course, course_id)
    if course is None:
        raise HTTPException(status_code=404, detail="Course not found")
    curriculum = db.get(Curriculum, course.curriculum_id)

    offerings = (
        db.query(CourseOffering)
        .filter(CourseOffering.course_id == course_id)
        .order_by(CourseOffering.academic_year, CourseOffering.semester, CourseOffering.section)
        .all()
    )
    instructor_ids = {o.instructor_id for o in offerings if o.instructor_id is not None}
    instructor_by_id = (
        {u.id: u for u in db.query(User).filter(User.id.in_(instructor_ids)).all()}
        if instructor_ids
        else {}
    )

    def instructor_name(offering: CourseOffering) -> str:
        instructor = instructor_by_id.get(offering.instructor_id)
        return _full_name(instructor.first_name, instructor.last_name) if instructor else "-"

    enrolled_count_by_offering: dict[int, int] = {}
    if offerings:
        for (offering_id,) in db.query(Enrollment.offering_id).filter(
            Enrollment.offering_id.in_([o.id for o in offerings])
        ):
            enrolled_count_by_offering[offering_id] = enrolled_count_by_offering.get(offering_id, 0) + 1

    clos = db.query(CLO).filter(CLO.course_id == course_id).order_by(CLO.code).all()

    # PLO ของรายวิชา = PLO ที่มี CLO ของวิชานี้ผูกอยู่ใน clo_plo_mapping
    plo_ids_by_clo: dict[int, set[int]] = {}
    plo_by_id: dict[int, PLO] = {}
    if clos:
        mapping_rows = (
            db.query(CLOPLOMapping.clo_id, PLO)
            .join(PLO, PLO.id == CLOPLOMapping.plo_id)
            .filter(CLOPLOMapping.clo_id.in_([c.id for c in clos]))
            .all()
        )
        for clo_id, plo in mapping_rows:
            plo_ids_by_clo.setdefault(clo_id, set()).add(plo.id)
            plo_by_id[plo.id] = plo
    course_plos = sorted(plo_by_id.values(), key=lambda p: p.code)

    workbook = Workbook()

    # --- ชีท 1: ข้อมูลรายวิชา ---
    ws = workbook.active
    ws.title = "ข้อมูลรายวิชา"
    ws.cell(row=1, column=1, value="ข้อมูลรายวิชา").font = _TITLE_FONT
    unique_instructors = list(dict.fromkeys(instructor_name(o) for o in offerings if o.instructor_id))
    row = _write_key_values(
        ws,
        3,
        [
            ("รหัสวิชา", course.course_code),
            ("ชื่อวิชา (ไทย)", course.name_th),
            ("ชื่อวิชา (อังกฤษ)", course.name_en or "-"),
            ("หน่วยกิต", course.credit),
            ("หมวดหมู่", course.category or "-"),
            ("หลักสูตร", f"{curriculum.name} ({curriculum.year})" if curriculum else "-"),
            ("อาจารย์ผู้สอน", ", ".join(unique_instructors) or "-"),
        ],
    )

    row += 1
    ws.cell(row=row, column=1, value="กลุ่มเรียนที่เปิดสอน").font = _SECTION_FONT
    row = _write_table(
        ws,
        row + 1,
        ["ปีการศึกษา/ภาคเรียน", "อาจารย์ผู้สอน", "จำนวนนักศึกษา"],
        [
            [_term_label(o), instructor_name(o), enrolled_count_by_offering.get(o.id, 0)]
            for o in offerings
        ],
    )

    row += 1
    ws.cell(row=row, column=1, value="CLO ของรายวิชา").font = _SECTION_FONT
    row = _write_table(
        ws,
        row + 1,
        ["รหัส CLO", "คำอธิบาย", "เกณฑ์ผ่าน (%)", "PLO ที่เกี่ยวข้อง"],
        [
            [
                clo.code,
                clo.description,
                float(clo.pass_threshold_percent),
                ", ".join(sorted(plo_by_id[pid].code for pid in plo_ids_by_clo.get(clo.id, ()))) or "-",
            ]
            for clo in clos
        ],
    )

    row += 1
    ws.cell(row=row, column=1, value="PLO ของรายวิชา").font = _SECTION_FONT
    _write_table(
        ws,
        row + 1,
        ["รหัส PLO", "คำอธิบาย", "หมวดหมู่", "CLO ที่เกี่ยวข้อง"],
        [
            [
                plo.code,
                plo.description_th,
                plo.category,
                ", ".join(c.code for c in clos if plo.id in plo_ids_by_clo.get(c.id, ())),
            ]
            for plo in course_plos
        ],
    )
    _set_column_widths(ws, [22, 60, 22, 22])

    # --- ชีท 2: ผล CLO นักศึกษา (1 แถวต่อ 1 การลงทะเบียน) ---
    ws2 = workbook.create_sheet("ผล CLO นักศึกษา")
    ws2.cell(row=1, column=1, value=f"ผลการผ่าน CLO — {course.course_code} {course.name_th}").font = _TITLE_FONT
    visible_offerings = [o for o in offerings if _can_see_offering_scores(o, current_user)]
    if len(visible_offerings) < len(offerings):
        ws2.cell(row=2, column=1, value=OWN_OFFERINGS_NOTE).font = _NOTE_FONT

    headers = ["ลำดับ", "รหัสนักศึกษา", "ชื่อ-นามสกุล", "รุ่น", "ภาคเรียนที่เรียน"]
    first_clo_col = len(headers)
    headers += [clo.code for clo in clos] + ["ผ่าน CLO"]

    enrollment_rows: list[tuple[Student, CourseOffering, list[str]]] = []
    for offering in visible_offerings:
        results = compute_offering_clo_achievement_raw(db, offering)
        status_by_student_clo: dict[tuple[str, int], str] = {}
        roster: dict[str, Student] = {}
        for result in results:
            for student_result in result.student_results:
                roster[student_result.student.id] = student_result.student
                if student_result.passed is None:
                    status = NO_DATA_TEXT
                else:
                    status = PASS_TEXT if student_result.passed else FAIL_TEXT
                status_by_student_clo[(student_result.student.id, result.clo.id)] = status
        if not results:
            # วิชายังไม่มี CLO — ยังต้องแสดงรายชื่อนักศึกษาที่ลงทะเบียนอยู่
            for student in (
                db.query(Student)
                .join(Enrollment, Enrollment.student_id == Student.id)
                .filter(Enrollment.offering_id == offering.id)
            ):
                roster[student.id] = student
        for student in roster.values():
            statuses = [status_by_student_clo.get((student.id, clo.id), NO_DATA_TEXT) for clo in clos]
            enrollment_rows.append((student, offering, statuses))
    enrollment_rows.sort(key=lambda r: (r[0].id, r[1].academic_year, r[1].semester, r[1].section))

    table_rows = [
        [
            index,
            student.id,
            _full_name(student.first_name, student.last_name, student.title),
            student.cohort_year,
            _term_label(offering),
            *statuses,
            f"{statuses.count(PASS_TEXT)}/{len(clos)}",
        ]
        for index, (student, offering, statuses) in enumerate(enrollment_rows, start=1)
    ]
    header_row = 4
    _write_table(
        ws2,
        header_row,
        headers,
        table_rows,
        status_columns=set(range(first_clo_col, first_clo_col + len(clos))),
    )
    ws2.freeze_panes = ws2.cell(row=header_row + 1, column=4)
    _set_column_widths(ws2, [7, 15, 30, 8, 18] + [11] * len(clos) + [11])

    return _xlsx_response(workbook, f"course_{course.course_code}.xlsx")


@router.get("/student/{student_id}")
def export_student_excel(
    student_id: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_role("admin", "instructor")),
):
    """
    ทำอะไร : สร้างไฟล์ Excel รายบุคคลของนักศึกษา 1 คน (2 ชีท — ดู docstring ของไฟล์)

    เชื่อมกับ : ผลบรรลุ PLO มาจาก _calculate_plo_achievement_for_student ตัวเดียวกับ GET
                /plo/achievement ที่หน้า /student-plo ใช้ (ค่า % ต่อเนื่อง + has_data) — "ผ่าน CLO" ต่อ
                วิชาคำนวณต่อกลุ่มเรียนด้วย compute_offering_clo_achievement_raw()

    ถ้าแก้ : ผลบรรลุ PLO เปิดให้ admin และอาจารย์ทุกคน (ข้อมูลระดับหลักสูตร) แต่คอลัมน์ "ผ่าน CLO" เป็น
             ข้อมูลระดับรายวิชา อาจารย์เห็นเฉพาะกลุ่มเรียนที่ตัวเองสอน (อื่นๆ แสดง "-") — 404 ถ้าไม่พบ
             นักศึกษา
    """
    student = db.get(Student, student_id)
    if student is None:
        raise HTTPException(status_code=404, detail="Student not found")
    curriculum = db.get(Curriculum, student.curriculum_id)

    plo_clo_weights, _, _ = _build_plo_requirements(db, student.curriculum_id)
    achievement = _calculate_plo_achievement_for_student(db, student, plo_clo_weights)
    plo_category_by_id = {
        plo.id: plo.category
        for plo in db.query(PLO).filter(PLO.curriculum_id == student.curriculum_id).all()
    }
    achieved_count = sum(1 for item in achievement.plo_achievements if item.is_achieved)

    enrollment_rows = (
        db.query(CourseOffering, Course)
        .join(Enrollment, Enrollment.offering_id == CourseOffering.id)
        .join(Course, Course.id == CourseOffering.course_id)
        .filter(Enrollment.student_id == student_id)
        .order_by(CourseOffering.academic_year, CourseOffering.semester, Course.course_code)
        .all()
    )
    instructor_ids = {o.instructor_id for o, _ in enrollment_rows if o.instructor_id is not None}
    instructor_by_id = (
        {u.id: u for u in db.query(User).filter(User.id.in_(instructor_ids)).all()}
        if instructor_ids
        else {}
    )

    workbook = Workbook()

    # --- ชีท 1: ข้อมูลนักศึกษา + ผลบรรลุ PLO ---
    ws = workbook.active
    ws.title = "ข้อมูลนักศึกษา"
    ws.cell(row=1, column=1, value="ข้อมูลนักศึกษา").font = _TITLE_FONT
    row = _write_key_values(
        ws,
        3,
        [
            ("รหัสนักศึกษา", student.id),
            ("ชื่อ-นามสกุล", _full_name(student.first_name, student.last_name, student.title)),
            ("รุ่น", student.cohort_year),
            ("ชั้นปี", student.current_year_level),
            ("สถานะ", student.status),
            ("หลักสูตร", f"{curriculum.name} ({curriculum.year})" if curriculum else "-"),
            ("บรรลุ PLO", f"{achieved_count}/{len(achievement.plo_achievements)} ข้อ"),
        ],
    )

    row += 1
    ws.cell(row=row, column=1, value="ผลการบรรลุ PLO").font = _SECTION_FONT
    plo_rows = []
    for item in achievement.plo_achievements:
        if not item.has_data:
            status = PLO_NO_DATA_TEXT
        elif item.is_achieved:
            status = PLO_ACHIEVED_TEXT
        else:
            status = PLO_NOT_ACHIEVED_TEXT
        plo_rows.append(
            [
                item.plo_code,
                item.description,
                plo_category_by_id.get(item.plo_id, "-"),
                round(item.achieved_percent, 1) if item.has_data else "-",
                status,
            ]
        )
    _write_table(
        ws,
        row + 1,
        ["รหัส PLO", "คำอธิบาย", "หมวดหมู่", "คะแนน PLO (%)", "ผลการบรรลุ"],
        plo_rows,
        status_columns={4},
    )
    _set_column_widths(ws, [18, 60, 18, 14, 16])

    # --- ชีท 2: รายวิชาที่เรียน ---
    ws2 = workbook.create_sheet("รายวิชาที่เรียน")
    ws2.cell(
        row=1,
        column=1,
        value=f"รายวิชาที่เรียน — {student.id} {_full_name(student.first_name, student.last_name, student.title)}",
    ).font = _TITLE_FONT

    course_rows = []
    hidden_any = False
    for index, (offering, course) in enumerate(enrollment_rows, start=1):
        if _can_see_offering_scores(offering, current_user):
            results = compute_offering_clo_achievement_raw(db, offering)
            passed = sum(
                1
                for result in results
                for student_result in result.student_results
                if student_result.student.id == student.id and student_result.passed
            )
            clo_summary = f"{passed}/{len(results)}"
        else:
            hidden_any = True
            clo_summary = "-"
        instructor = instructor_by_id.get(offering.instructor_id)
        course_rows.append(
            [
                index,
                _term_label(offering),
                course.course_code,
                course.name_th,
                course.credit,
                _full_name(instructor.first_name, instructor.last_name) if instructor else "-",
                clo_summary,
            ]
        )
    if hidden_any:
        ws2.cell(row=2, column=1, value=OWN_OFFERINGS_NOTE).font = _NOTE_FONT
    _write_table(
        ws2,
        4,
        ["ลำดับ", "ปีการศึกษา/ภาคเรียน", "รหัสวิชา", "ชื่อวิชา", "หน่วยกิต", "อาจารย์ผู้สอน", "ผ่าน CLO"],
        course_rows,
    )
    ws2.freeze_panes = "A5"
    _set_column_widths(ws2, [7, 18, 14, 40, 9, 28, 11])

    return _xlsx_response(workbook, f"student_{student.id}.xlsx")
