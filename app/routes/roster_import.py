"""API routes for importing official university Excel roster files (.xls/.xlsx).

ไฟล์ที่มหาวิทยาลัยส่งให้อาจารย์เป็นไฟล์ Excel รูปแบบเก่า (.xls, binary "CDFV2 Microsoft Excel"
ไม่ใช่ .xlsx) มีโครงสร้างคงที่: หัวเรื่องมหาวิทยาลัย, บรรทัด "ภาคการศึกษา n/pppp", บรรทัด
"รหัสวิชา <code> <ชื่อวิชา> หน่วยกิต <n> Sec <section>", บรรทัด "ผู้สอน <ชื่อ1>[, <ชื่อ2>, ...]",
ตามด้วยตารางรายชื่อที่มีหัวคอลัมน์ "เลขที่ / รหัสประจำตัว / ชื่อ / คะแนน / สัปดาห์..."

แทนที่จะอิงตำแหน่งแถว/คอลัมน์ตายตัว (ซึ่งขยับได้ 1 แถวถ้า "วันเวลาเรียน" มี 1 หรือ 2 บรรทัด) โค้ดนี้
สแกนหาข้อความ/หัวคอลัมน์ที่รู้จักในทุกเซลล์แทน เพื่อให้ทนทานต่อความต่างเล็กน้อยระหว่างไฟล์แต่ละรุ่น
(ตรวจสอบกับไฟล์จริง 5 ไฟล์ คนละรุ่น/section/ภาคเรียนแล้ว)

หลักการ import (ตามที่ผู้ใช้ยืนยัน):
- ไฟล์คือความจริง (authoritative) — ถ้านักศึกษามีอยู่แล้วแต่ชื่อในไฟล์ไม่ตรงกับในระบบ ให้แก้ตามไฟล์
- ถ้ายังไม่มี course_offering ตรงกับวิชา/ปีการศึกษา/ภาคเรียน/section ในไฟล์ ให้สร้างให้อัตโนมัติ
- ถ้าชื่อผู้สอนในไฟล์ไม่ตรงกับ user ในระบบ ให้สร้างบัญชีอาจารย์จริงให้เลย (username รูปแบบ
  "ajarn<id>" + รหัสผ่านชั่วคราวสุ่ม แสดงครั้งเดียวตอนสร้าง - เหมือน ajarn.somsak/ajarn.suda ที่มีอยู่
  แต่ไม่มีวิธี transliterate ชื่อไทยเป็นอังกฤษได้แม่นยำอัตโนมัติ จึงใช้ id เป็นส่วนต่อท้ายแทน)
- รายวิชา (Course) ต้องมีอยู่ในระบบก่อนแล้วเท่านั้น (import จะไม่สร้างรายวิชาใหม่) - ถ้าไม่เจอจะแจ้ง error

รองรับ 2 โหมดในการเรียกเดียวกัน ผ่าน form field `dry_run`:
- `dry_run=true` (ค่าเริ่มต้น): parse + คำนวณผลลัพธ์ที่ *จะ* เกิดขึ้น แต่ rollback ไม่บันทึกจริง
  ใช้แสดง preview ให้ผู้ใช้ตรวจสอบก่อน
- `dry_run=false`: บันทึกจริงทั้งหมด (สร้าง/แก้ course_offering, user, student, enrollment)
"""
from __future__ import annotations

import io
import re
import secrets
from collections import Counter
from dataclasses import dataclass
from datetime import date

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from openpyxl import load_workbook
from sqlalchemy.dialects.postgresql import insert as pg_insert
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.auth import hash_password, require_role
from app.database import get_db
from app.models import Course, CourseOffering, Curriculum, Enrollment, Student, User
from app.schemas import RosterImportInstructor, RosterImportResponse, RosterImportStudentRow

router = APIRouter(prefix="/roster-import", tags=["Roster Import"])

STUDENT_TITLES = ["นางสาว", "นาย", "นาง"]  # เช็ค "นางสาว" ก่อน "นาง" เพราะ "นาง" เป็นคำนำหน้าของ "นางสาว"
ENGLISH_TITLE_RE = re.compile(r"^(Mr|Mrs|Ms|Miss)\.?\s*", re.IGNORECASE)
SEMESTER_RE = re.compile(r"ภาคการศึกษา\s*(\d+)\s*/\s*(\d+)")
COURSE_LINE_RE = re.compile(r"รหัสวิชา\s+(\S+)\s+(.+?)\s+หน่วยกิต\s+\S+\s+Sec\s+(\S+)")


class RosterParseError(Exception):
    """ไฟล์ไม่ตรงรูปแบบที่คาดไว้ - ส่งเป็น HTTP 400 กลับให้ผู้ใช้"""


@dataclass
class ParsedRoster:
    course_code: str
    course_name: str
    section: str
    academic_year: int
    semester: int
    cohort_year: int | None
    instructor_names: list[str]
    students: list[tuple[str, str | None, str, str]]  # (student_id, title, first_name, last_name)


def _cellstr(value) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _load_xls_workbook(content: bytes):
    """xlrd's own OLE2 (compound document) reader is stricter than real Excel about
    the sector-chain layout, and some genuine files (observed from the university's
    export tool) fail with `IndexError` inside xlrd's compdoc even though Excel opens
    them fine and `olefile` (a more tolerant standalone OLE2 reader) reads them with
    zero issues. Fall back to extracting the raw 'Workbook'/'Book' BIFF stream with
    olefile and handing those bytes to xlrd directly - xlrd accepts a raw BIFF stream
    (no OLE2 signature) and skips its own compdoc parsing entirely in that case."""
    import xlrd

    try:
        return xlrd.open_workbook(file_contents=content)
    except Exception:
        pass

    try:
        import olefile
    except ImportError as exc:  # pragma: no cover - ต้องติดตั้ง olefile บนเซิร์ฟเวอร์จริง
        raise RosterParseError(
            "ไฟล์ .xls นี้เปิดด้วยวิธีปกติไม่ได้ และเซิร์ฟเวอร์ไม่มีไลบรารีสำรอง (olefile) - กรุณาติดต่อผู้ดูแลระบบ"
        ) from exc

    try:
        ole = olefile.OleFileIO(io.BytesIO(content))
        try:
            stream_name = "Workbook" if ole.exists("Workbook") else "Book"
            stream_bytes = ole.openstream(stream_name).read()
        finally:
            ole.close()
        return xlrd.open_workbook(file_contents=stream_bytes)
    except Exception as exc:
        raise RosterParseError(
            "ไฟล์ .xls นี้เปิดไม่ได้ (โครงสร้างไฟล์เสียหายหรือไม่ตรงรูปแบบ) - "
            "ลองเปิดไฟล์นี้ด้วย Excel แล้ว \"บันทึกเป็น\" ใหม่ (.xls หรือ .xlsx) แล้วอัปโหลดไฟล์ที่บันทึกใหม่แทน"
        ) from exc


def _load_rows(filename: str, content: bytes) -> list[list]:
    lower_name = filename.lower()
    if lower_name.endswith(".xls"):
        try:
            import xlrd  # noqa: F401  (สำหรับเช็คว่าติดตั้งไว้หรือยัง ก่อนเข้า _load_xls_workbook)
        except ImportError as exc:  # pragma: no cover - ต้องติดตั้ง xlrd บนเซิร์ฟเวอร์จริง
            raise RosterParseError(
                "เซิร์ฟเวอร์ยังไม่ได้ติดตั้งไลบรารีสำหรับอ่านไฟล์ .xls (xlrd) - กรุณาติดต่อผู้ดูแลระบบ"
            ) from exc
        book = _load_xls_workbook(content)
        sheet = book.sheet_by_index(0)
        return [[sheet.cell_value(r, c) for c in range(sheet.ncols)] for r in range(sheet.nrows)]
    if lower_name.endswith(".xlsx"):
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        worksheet = workbook.active
        return [list(row) for row in worksheet.iter_rows(values_only=True)]
    raise RosterParseError("รองรับเฉพาะไฟล์ .xls หรือ .xlsx เท่านั้น (ไฟล์ที่มหาวิทยาลัยส่งให้ปกติเป็น .xls)")


def split_student_name(raw: str) -> tuple[str | None, str, str]:
    """แยกคำนำหน้า/ชื่อ/นามสกุลจากชื่อนักศึกษา เช่น 'นายกฤษณะ  เข็มมา' หรือ 'Mr.Socheat  Moeun'
    ชื่อ-นามสกุลคั่นด้วยช่องว่าง 2 ตัวขึ้นไปเสมอ (ยืนยันจากไฟล์จริง 5 ไฟล์)"""
    s = raw.strip()
    title: str | None = None
    for t in STUDENT_TITLES:
        if s.startswith(t):
            title = t
            s = s[len(t):]
            break
    if title is None:
        m = ENGLISH_TITLE_RE.match(s)
        if m:
            title = m.group(1) + "."
            s = s[m.end():]
    s = s.strip()
    parts = re.split(r"\s{2,}", s)
    if len(parts) >= 2 and parts[0] and parts[1]:
        first, last = parts[0].strip(), " ".join(parts[1:]).strip()
    else:
        tokens = s.split()
        if len(tokens) >= 2:
            first, last = tokens[0], " ".join(tokens[1:])
        else:
            first, last = s, ""
    return title, first, last


def split_instructor_name(raw: str) -> tuple[str | None, str, str]:
    """แยกคำนำหน้า/ชื่อ/นามสกุลจากชื่อผู้สอน เช่น 'อาจารย์ ชลัท รังสิมาเทวัญ' หรือ
    'ผู้ช่วยศาสตราจารย์ ดร. สมศักดิ์ จีวัฒนา' - คั่นด้วยช่องว่างเดียว 2 คำสุดท้ายคือชื่อ-นามสกุล
    ที่เหลือด้านหน้าคือคำนำหน้า (รองรับคำนำหน้ายศวิชาการยาวกี่คำก็ได้)"""
    tokens = raw.strip().split()
    if len(tokens) >= 3:
        first, last = tokens[-2], tokens[-1]
        title = " ".join(tokens[:-2])
    elif len(tokens) == 2:
        title, first, last = None, tokens[0], tokens[1]
    elif len(tokens) == 1:
        title, first, last = None, tokens[0], ""
    else:
        title, first, last = None, "", ""
    return title, first, last


def compute_year_level(cohort_year_2digit: int) -> int:
    """คำนวณชั้นปีปัจจุบันของนักศึกษาใหม่ ณ วันที่รันจริง (ไม่ใช่ตอน parse ไฟล์) จากปี พ.ศ. 2 หลักของรุ่น
    เช่น cohort_year=66 → เข้าปีการศึกษา 2566 ปีการศึกษาเริ่มประมาณเดือนมิถุนายน จึงนับปีการศึกษาปัจจุบัน
    ถอยหลัง 1 ถ้ายังไม่ถึงเดือนมิถุนายน แล้ว clamp ผลลัพธ์ไว้ระหว่างปี 1-4"""
    today = date.today()
    buddhist_2digit = (today.year + 543) % 100
    if today.month < 6:
        buddhist_2digit = (buddhist_2digit - 1) % 100
    level = buddhist_2digit - cohort_year_2digit + 1
    return max(1, min(4, level))


def parse_roster_xls(filename: str, content: bytes) -> ParsedRoster:
    rows = _load_rows(filename, content)
    rows = [row for row in rows if row]
    if not rows:
        raise RosterParseError("ไฟล์ว่างเปล่าหรืออ่านไม่ได้")

    semester = academic_year = None
    for row in rows:
        for cell in row:
            m = SEMESTER_RE.search(_cellstr(cell))
            if m:
                semester, academic_year = int(m.group(1)), int(m.group(2))
                break
        if semester is not None:
            break
    if semester is None:
        raise RosterParseError("ไม่พบข้อมูลภาคการศึกษา/ปีการศึกษาในไฟล์ (ต้องมีข้อความรูปแบบ 'ภาคการศึกษา n/pppp')")

    course_code = course_name = section = None
    course_row_idx = None
    for idx, row in enumerate(rows):
        for cell in row:
            m = COURSE_LINE_RE.search(_cellstr(cell))
            if m:
                course_code, course_name, section = m.group(1).strip(), m.group(2).strip(), m.group(3).strip()
                course_row_idx = idx
                break
        if course_code is not None:
            break
    if course_code is None:
        raise RosterParseError("ไม่พบข้อมูลรหัสวิชา/Section ในไฟล์ (ต้องมีข้อความขึ้นต้นด้วย 'รหัสวิชา')")

    instructor_names: list[str] = []
    search_rows = rows[course_row_idx:course_row_idx + 4] if course_row_idx is not None else rows
    for row in search_rows:
        for cell in row:
            s = _cellstr(cell)
            if s.startswith("ผู้สอน"):
                raw = s[len("ผู้สอน"):].strip()
                instructor_names = [p.strip() for p in raw.split(",") if p.strip()]
                break
        if instructor_names:
            break

    header_row_idx = None
    id_col = name_col = None
    for idx, row in enumerate(rows):
        cells = [_cellstr(c) for c in row]
        if "เลขที่" in cells and "รหัสประจำตัว" in cells:
            header_row_idx = idx
            id_col = cells.index("รหัสประจำตัว")
            name_col = cells.index("ชื่อ") if "ชื่อ" in cells else None
            break
    if header_row_idx is None or name_col is None:
        raise RosterParseError("ไม่พบหัวตารางรายชื่อนักศึกษา (คอลัมน์ 'เลขที่ / รหัสประจำตัว / ชื่อ')")

    students: list[tuple[str, str | None, str, str]] = []
    seen_ids: set[str] = set()
    started = False
    for row in rows[header_row_idx + 1:]:
        sid = _cellstr(row[id_col]) if id_col < len(row) else ""
        if sid.isdigit() and len(sid) >= 8:
            started = True
            if sid in seen_ids:
                continue
            seen_ids.add(sid)
            name_raw = _cellstr(row[name_col]) if name_col < len(row) else ""
            title, first, last = split_student_name(name_raw)
            students.append((sid, title, first, last))
        elif started:
            break

    if not students:
        raise RosterParseError("ไม่พบรายชื่อนักศึกษาในไฟล์")

    prefixes = [sid[:2] for sid, *_ in students if sid[:2].isdigit()]
    cohort_year = int(Counter(prefixes).most_common(1)[0][0]) if prefixes else None

    return ParsedRoster(
        course_code=course_code,
        course_name=course_name,
        section=section,
        academic_year=academic_year,
        semester=semester,
        cohort_year=cohort_year,
        instructor_names=instructor_names,
        students=students,
    )


def _apply_roster_import(db: Session, parsed: ParsedRoster, commit: bool) -> RosterImportResponse:
    errors: list[str] = []

    courses = (
        db.query(Course)
        .join(Curriculum)
        .filter(Course.course_code == parsed.course_code)
        .order_by(Curriculum.is_active.desc(), Course.id)
        .all()
    )
    if not courses:
        errors.append(
            f"ไม่พบวิชารหัส {parsed.course_code} ในระบบ - ระบบนี้ไม่สร้างรายวิชาใหม่อัตโนมัติ "
            f"กรุณาสร้างวิชา '{parsed.course_code} {parsed.course_name}' ก่อน แล้วนำเข้าไฟล์นี้ใหม่อีกครั้ง"
        )
        return RosterImportResponse(
            dry_run=not commit,
            course_code=parsed.course_code,
            course_name_th=parsed.course_name,
            course_found=False,
            academic_year=parsed.academic_year,
            semester=parsed.semester,
            section=parsed.section,
            cohort_year=parsed.cohort_year,
            offering_action="error",
            errors=errors,
        )
    course = courses[0]
    if len(courses) > 1:
        errors.append(
            f"พบวิชารหัส {parsed.course_code} มากกว่า 1 หลักสูตรในระบบ - เลือกหลักสูตร "
            f"'{course.curriculum.name}' (id={course.curriculum_id}) ให้อัตโนมัติเพราะเป็นหลักสูตรที่ active อยู่"
        )

    # --- ผู้สอน ---
    instructor_results: list[RosterImportInstructor] = []
    primary_instructor_id: int | None = None
    created_credentials: list[dict] = []
    for idx, raw_name in enumerate(parsed.instructor_names):
        title, first, last = split_instructor_name(raw_name)
        existing = db.query(User).filter(User.first_name == first, User.last_name == last).first()
        if existing:
            entry = RosterImportInstructor(
                full_name=raw_name, title=title, action="matched_existing",
                user_id=existing.id, username=existing.username,
            )
            if idx == 0:
                primary_instructor_id = existing.id
        elif commit:
            temp_password = secrets.token_urlsafe(9)
            new_user = User(
                username="__pending__",
                password=hash_password(temp_password),
                first_name=first,
                last_name=last,
                role="instructor",
            )
            db.add(new_user)
            db.flush()
            new_user.username = f"ajarn{new_user.id}"
            db.flush()
            entry = RosterImportInstructor(
                full_name=raw_name, title=title, action="created",
                user_id=new_user.id, username=new_user.username, temp_password=temp_password,
            )
            created_credentials.append({
                "username": new_user.username,
                "temp_password": temp_password,
                "full_name": f"{(title + ' ') if title else ''}{first} {last}".strip(),
            })
            if idx == 0:
                primary_instructor_id = new_user.id
        else:
            entry = RosterImportInstructor(full_name=raw_name, title=title, action="will_create")
        instructor_results.append(entry)

    # --- course offering ---
    offering = (
        db.query(CourseOffering)
        .filter_by(
            course_id=course.id,
            academic_year=parsed.academic_year,
            semester=parsed.semester,
            section=parsed.section,
        )
        .first()
    )
    offering_id: int | None
    if offering:
        offering_action = "matched_existing"
        offering_id = offering.id
    elif commit:
        offering = CourseOffering(
            course_id=course.id,
            instructor_id=primary_instructor_id,
            cohort_year=parsed.cohort_year,
            academic_year=parsed.academic_year,
            semester=parsed.semester,
            section=parsed.section,
        )
        db.add(offering)
        db.flush()
        offering_action = "created"
        offering_id = offering.id
    else:
        offering_action = "will_create"
        offering_id = None

    # --- นักศึกษา ---
    student_rows: list[RosterImportStudentRow] = []
    to_enroll_ids: list[str] = []
    counts = {"create": 0, "update_name": 0, "unchanged": 0, "error": 0}
    for i, (sid, title, first, last) in enumerate(parsed.students, start=1):
        existing_student = db.get(Student, sid)
        if existing_student is None:
            level = compute_year_level(parsed.cohort_year if parsed.cohort_year is not None else int(sid[:2]))
            row = RosterImportStudentRow(
                line_no=i, student_id=sid, title=title, first_name=first, last_name=last, action="create",
            )
            counts["create"] += 1
            if commit:
                db.add(Student(
                    id=sid,
                    curriculum_id=course.curriculum_id,
                    first_name=first,
                    last_name=last,
                    title=title,
                    cohort_year=parsed.cohort_year if parsed.cohort_year is not None else int(sid[:2]),
                    current_year_level=level,
                ))
            to_enroll_ids.append(sid)
        elif existing_student.curriculum_id != course.curriculum_id:
            row = RosterImportStudentRow(
                line_no=i, student_id=sid, title=title, first_name=first, last_name=last, action="error",
                detail=(
                    f"นักศึกษาคนนี้อยู่คนละหลักสูตรกับวิชานี้ในระบบ (curriculum_id="
                    f"{existing_student.curriculum_id} ไม่ตรงกับ {course.curriculum_id}) - ข้ามการลงทะเบียนให้ "
                    "กรุณาตรวจสอบด้วยตนเอง"
                ),
            )
            counts["error"] += 1
        else:
            name_changed = (
                existing_student.first_name != first
                or existing_student.last_name != last
                or (existing_student.title or None) != (title or None)
            )
            if name_changed:
                row = RosterImportStudentRow(
                    line_no=i, student_id=sid, title=title, first_name=first, last_name=last,
                    action="update_name",
                    detail=(
                        f"ชื่อเดิมในระบบ: {existing_student.title or ''} {existing_student.first_name} "
                        f"{existing_student.last_name} -> ในไฟล์: {title or ''} {first} {last}"
                    ),
                )
                counts["update_name"] += 1
                if commit:
                    existing_student.first_name = first
                    existing_student.last_name = last
                    existing_student.title = title
            else:
                row = RosterImportStudentRow(
                    line_no=i, student_id=sid, title=title, first_name=first, last_name=last, action="unchanged",
                )
                counts["unchanged"] += 1
            to_enroll_ids.append(sid)
        student_rows.append(row)

    if commit:
        # นักศึกษาใหม่ที่เพิ่ง db.add() ไว้ข้างบนต้อง flush ลง DB ก่อน ไม่งั้น bulk insert
        # enrollment ด้านล่าง (Core insert ตรง ไม่ผ่าน ORM flush อัตโนมัติ) จะชน FK constraint
        # เพราะ student row ยังไม่มีอยู่จริงในตาราง
        db.flush()

    # --- ลงทะเบียนเรียน ---
    enrollments_added = 0
    enrollments_already = 0
    if to_enroll_ids:
        if commit and offering_id is not None:
            already = {
                r[0]
                for r in db.query(Enrollment.student_id)
                .filter(Enrollment.offering_id == offering_id, Enrollment.student_id.in_(to_enroll_ids))
                .all()
            }
            to_insert = [sid for sid in to_enroll_ids if sid not in already]
            if to_insert:
                stmt = (
                    pg_insert(Enrollment.__table__)
                    .values([{"student_id": sid, "offering_id": offering_id} for sid in to_insert])
                    .on_conflict_do_nothing(index_elements=["student_id", "offering_id"])
                    .returning(Enrollment.__table__.c.student_id)
                )
                result = db.execute(stmt)
                enrollments_added = len(result.fetchall())
            enrollments_already = len(already)
        elif offering_id is not None:
            already = {
                r[0]
                for r in db.query(Enrollment.student_id)
                .filter(Enrollment.offering_id == offering_id, Enrollment.student_id.in_(to_enroll_ids))
                .all()
            }
            enrollments_already = len(already)
            enrollments_added = len(to_enroll_ids) - len(already)
        else:
            # dry-run กับ course_offering ที่ยังไม่มีจริง - ยังไม่มีใครลงทะเบียนได้เลย
            enrollments_added = len(to_enroll_ids)

    if commit:
        try:
            db.commit()
        except IntegrityError as exc:
            db.rollback()
            raise HTTPException(
                status_code=409, detail=f"นำเข้าไม่สำเร็จ - ข้อมูลขัดแย้งกันในฐานข้อมูล ({exc.orig})"
            ) from exc
        if offering is not None:
            db.refresh(offering)
            offering_id = offering.id
    else:
        db.rollback()

    summary = {
        "students_total": len(parsed.students),
        "students_create": counts["create"],
        "students_update_name": counts["update_name"],
        "students_unchanged": counts["unchanged"],
        "students_error": counts["error"],
        "enrollments_added": enrollments_added,
        "enrollments_already": enrollments_already,
        "instructors_created": sum(1 for r in instructor_results if r.action == "created"),
        "instructors_will_create": sum(1 for r in instructor_results if r.action == "will_create"),
    }

    return RosterImportResponse(
        dry_run=not commit,
        course_code=parsed.course_code,
        course_name_th=course.name_th,
        course_found=True,
        academic_year=parsed.academic_year,
        semester=parsed.semester,
        section=parsed.section,
        cohort_year=parsed.cohort_year,
        offering_id=offering_id,
        offering_action=offering_action,
        instructors=instructor_results,
        students=student_rows,
        enrollments_added=enrollments_added,
        enrollments_already=enrollments_already,
        summary=summary,
        new_instructor_credentials=created_credentials,
        errors=errors,
    )


@router.post("", response_model=RosterImportResponse)
async def import_roster(
    file: UploadFile = File(...),
    dry_run: bool = Form(True),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_role("admin")),
):
    """นำเข้าไฟล์รายชื่อจากมหาวิทยาลัย (.xls/.xlsx) - admin เท่านั้น เพราะ endpoint นี้สร้างบัญชี
    ผู้ใช้ (อาจารย์) ใหม่ในระบบได้ ซึ่งเป็นสิทธิ์ระดับแอดมินทุกจุดในระบบนี้อยู่แล้ว (ดู users.py)
    เรียกด้วย dry_run=true ก่อนเพื่อดูตัวอย่างผลลัพธ์ แล้วค่อยเรียกซ้ำด้วย dry_run=false เพื่อบันทึกจริง"""
    content = await file.read()
    try:
        parsed = parse_roster_xls(file.filename or "", content)
    except RosterParseError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc

    return _apply_roster_import(db, parsed, commit=not dry_run)
