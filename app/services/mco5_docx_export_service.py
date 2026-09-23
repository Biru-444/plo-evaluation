"""
ทำอะไร : สร้างไฟล์ Word (.docx) ประกอบการกรอก มคอ.5 ตามโครงแบบฟอร์ม OBE5 BRU (Phase 2 - ดู
         TASK-export-mco5.md ข้อ 6) - หัวเรื่อง "รายงานผลการดำเนินการของรายวิชา ตามเกณฑ์คุณภาพหลักสูตร
         แบบมุ่งเน้นผลลัพธ์การเรียนรู้" หมวด 1-6 + ส่วนลงนาม เติมข้อมูลอัตโนมัติเฉพาะส่วนที่ระบบมีข้อมูล
         จริง (หมวด 1, ตาราง CLO หมวด 2 ข้อ 3, หมวด 3 ข้อ 1-4 และ 7) ส่วนที่เหลือใส่
         "[อาจารย์ผู้สอนกรอก]" ตัวเอียงสีเทาไว้แทน - ไม่มีข้อมูลรายบุคคล (ต่างจาก Excel ที่มีชีต 5)

เชื่อมกับ : ใช้ app/services/mco5_data_service.py ชุดเดียวกับที่ Excel export
            (mco5_export_service.py) ใช้ทุกฟังก์ชัน (compute_course_info_rows, compute_clo_rows,
            compute_grade_distribution, compute_assessment_confirmation_rows) - ไม่คำนวณ/query ซ้ำเอง
            เลยสักจุด ต่างกันแค่ชั้นวาดผล (python-docx แทน openpyxl) - เรียกจาก
            app/routes/clo_calculation.py::export_mco5_docx

ถ้าแก้ : ฟอนต์ต้องเป็น "TH Sarabun New" เสมอ (ไฟล์นี้จะถูกคัดลอกไปใส่แบบฟอร์มราชการจริง) - ตัวอักษรไทยใน
         Word ต้องตั้งทั้ง ascii/hAnsi (_set_font ตั้งให้ผ่าน run.font.name) และ w:cs (complex script -
         Word ใช้ font slot นี้กับอักษรไทยจริง ๆ ไม่ใช่ eastAsia ที่เป็นของ CJK) ไม่งั้น Word อาจ fallback
         ไปฟอนต์ default เงียบๆ ทั้งที่ตั้ง run.font.name ถูกแล้ว
"""
from __future__ import annotations

from decimal import Decimal
from io import BytesIO

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml.shared import OxmlElement
from docx.shared import Pt, RGBColor
from sqlalchemy.orm import Session

from app.models import CourseOffering
from app.services.clo_achievement_service import compute_offering_clo_achievement_raw
from app.services.mco5_data_service import (
    AssessmentConfirmationRow,
    CLORow,
    CourseInfoRow,
    GradeDistribution,
    compute_assessment_confirmation_rows,
    compute_clo_rows,
    compute_course_info_rows,
    compute_grade_distribution,
)

FONT_NAME = "TH Sarabun New"
BODY_SIZE = Pt(16)
HEADING_SIZE = Pt(18)
PLACEHOLDER_TEXT = "[อาจารย์ผู้สอนกรอก]"
GRAY = RGBColor(0x80, 0x80, 0x80)


def _apply_complex_script_properties(rpr, size: Pt, *, bold: bool = False, italic: bool = False) -> None:
    """ตั้ง font name/ขนาด/ตัวหนา/ตัวเอียงสำหรับ "complex script" (w:cs, w:szCs, w:bCs, w:iCs) ให้ rPr
    ที่ส่งมา (ใช้ได้ทั้ง run-level และ style-level - โครง XML เหมือนกัน) - Word แยกอักษรไทยเป็นคนละ slot
    จาก western text (w:rFonts ascii/hAnsi, w:sz, w:b, w:i ที่ python-docx's run.font.*/style.font.*
    ตั้งให้อัตโนมัติอยู่แล้ว) ถ้าไม่ตั้ง cs คู่กันให้ครบทั้ง 4 ตัวตรงนี้ Word จะ fallback ไปค่า default
    (มักเป็น 10pt ไม่ตัวหนา) สำหรับอักษรไทยโดยเฉพาะ แม้ font name (w:cs) จะตั้งถูกแล้วก็ตาม - เคยพลาดจุด
    szCs/bCs/iCs มาก่อน (แก้เฉพาะ w:cs อย่างเดียว) หัวข้อหมวดที่ควรตัวหนา 18pt เลยจะแสดงเป็นค่า default
    แทนสำหรับตัวอักษรไทย"""
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:cs"), FONT_NAME)

    szcs = rpr.find(qn("w:szCs"))
    if szcs is None:
        szcs = OxmlElement("w:szCs")
        rpr.append(szcs)
    szcs.set(qn("w:val"), str(int(size.pt * 2)))  # หน่วยเป็นครึ่งจุด (half-points) เหมือน w:sz

    if bold and rpr.find(qn("w:bCs")) is None:
        rpr.append(OxmlElement("w:bCs"))
    if italic and rpr.find(qn("w:iCs")) is None:
        rpr.append(OxmlElement("w:iCs"))


def _set_font(run, *, size: Pt = BODY_SIZE, bold: bool = False, italic: bool = False, color=None) -> None:
    """ตั้งฟอนต์ TH Sarabun New ให้ run - ตั้งทั้ง ascii/hAnsi ผ่าน run.font.name ปกติ (western) และ
    complex-script properties ทั้งชุด (w:cs/w:szCs/w:bCs/w:iCs) ผ่าน _apply_complex_script_properties
    เพราะ python-docx ไม่มี property สำหรับฝั่ง complex script โดยตรงเลย"""
    run.font.name = FONT_NAME
    run.font.size = size
    run.font.bold = bold
    run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color
    _apply_complex_script_properties(run._element.get_or_add_rPr(), size, bold=bold, italic=italic)


def _add_paragraph(doc: Document, text: str, *, bold: bool = False, italic: bool = False, gray: bool = False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    _set_font(run, bold=bold, italic=italic, color=GRAY if gray else None)
    return p


def _add_placeholder(doc: Document) -> None:
    _add_paragraph(doc, PLACEHOLDER_TEXT, italic=True, gray=True)


def _add_category_heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    run = p.add_run(text)
    _set_font(run, size=HEADING_SIZE, bold=True)


def _add_subitem(doc: Document, text: str) -> None:
    _add_paragraph(doc, text, bold=True)


def _style_table(table) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"


def _set_cell_text(cell, text: str, *, bold: bool = False, header: bool = False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    if header:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    _set_font(run, bold=bold or header)


def _add_header_obe5_label(doc: Document) -> None:
    """header มุมขวาบนของทุกหน้า - 'OBE5 BRU' ตามสเปก"""
    header = doc.sections[0].header
    p = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("OBE5 BRU")
    _set_font(run, bold=True)


def _add_title(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(
        "รายงานผลการดำเนินการของรายวิชา\nตามเกณฑ์คุณภาพหลักสูตรแบบมุ่งเน้นผลลัพธ์การเรียนรู้"
    )
    _set_font(run, size=HEADING_SIZE, bold=True)


def _add_course_info_table(doc: Document, rows: list[CourseInfoRow]) -> None:
    table = doc.add_table(rows=len(rows), cols=2)
    _style_table(table)
    for i, row in enumerate(rows):
        _set_cell_text(table.cell(i, 0), row.label, bold=True)
        _set_cell_text(table.cell(i, 1), row.value)


def _add_clo_table(doc: Document, clo_rows: list[CLORow]) -> None:
    """ตาราง หมวด 2 ข้อ 3 - 5 คอลัมน์ตามสเปก: CLOs | กลยุทธ์การสอน (เว้นว่างเสมอ - ระบบไม่มีข้อมูลนี้) |
    วิธีการประเมินผล | ผลที่เกิดกับนักศึกษา | แนวทางพัฒนา"""
    headers = ["CLOs", "กลยุทธ์การสอน", "วิธีการประเมินผล", "ผลที่เกิดกับนักศึกษา", "แนวทางพัฒนา"]
    table = doc.add_table(rows=1, cols=len(headers))
    _style_table(table)
    for col, header in enumerate(headers):
        _set_cell_text(table.cell(0, col), header, header=True)

    for row in clo_rows:
        cells = table.add_row().cells
        _set_cell_text(cells[0], f"{row.clo_code}: {row.description}")
        _set_cell_text(cells[1], "")  # กลยุทธ์การสอน - ระบบไม่มีข้อมูลนี้ เว้นว่างตามสเปก (ไม่ใช่ placeholder)
        _set_cell_text(cells[2], row.assessment_text)
        _set_cell_text(cells[3], row.outcome_text)
        _set_cell_text(
            cells[4], row.improvement_text if row.improvement_text else "-"
        )


def _add_grade_table(doc: Document, dist: GradeDistribution) -> None:
    headers = ["ระดับคะแนน", "ความหมาย", "จำนวน (คน)", "ร้อยละ"]
    table = doc.add_table(rows=1, cols=len(headers))
    _style_table(table)
    for col, header in enumerate(headers):
        _set_cell_text(table.cell(0, col), header, header=True)

    for row in dist.rows:
        cells = table.add_row().cells
        _set_cell_text(cells[0], row.grade)
        _set_cell_text(cells[1], row.meaning)
        _set_cell_text(cells[2], str(row.count))
        _set_cell_text(cells[3], row.percent_display)

    cells = table.add_row().cells
    _set_cell_text(cells[0], "ยังไม่มีเกรด")
    _set_cell_text(cells[1], "-")
    _set_cell_text(cells[2], str(dist.no_grade_count))
    _set_cell_text(cells[3], dist.no_grade_percent_display)

    cells = table.add_row().cells
    _set_cell_text(cells[0], "รวม", bold=True)
    _set_cell_text(cells[2], str(dist.total_registered), bold=True)
    _set_cell_text(cells[3], dist.total_percent_display, bold=True)


def _add_assessment_table(doc: Document, rows: list[AssessmentConfirmationRow]) -> None:
    headers = ["วิธีการ", "CLO ที่วัด", "คะแนนเต็ม", "คะแนนเฉลี่ย (%)", "จำนวนคนที่มีคะแนน", "สรุปผล"]
    table = doc.add_table(rows=1, cols=len(headers))
    _style_table(table)
    for col, header in enumerate(headers):
        _set_cell_text(table.cell(0, col), header, header=True)

    for row in rows:
        cells = table.add_row().cells
        _set_cell_text(cells[0], row.method_label)
        _set_cell_text(cells[1], row.clo_codes_label)
        _set_cell_text(cells[2], str(row.total_score))
        avg_text = f"{row.average_percent:.1f}" if row.average_percent is not None else "-"
        _set_cell_text(cells[3], avg_text)
        _set_cell_text(cells[4], str(row.count_with_scores))
        _set_cell_text(cells[5], row.summary)


def _add_signature_section(doc: Document, offering: CourseOffering) -> None:
    doc.add_paragraph()
    instructor_name = (
        f"{offering.instructor.first_name} {offering.instructor.last_name}"
        if offering.instructor is not None
        else PLACEHOLDER_TEXT
    )
    _add_paragraph(doc, "ลงชื่อ .................................................... อาจารย์ผู้สอน")
    _add_paragraph(doc, f"({instructor_name})")


def build_mco5_docx(db: Session, offering: CourseOffering, target_rate: Decimal) -> BytesIO:
    """ทำอะไร : สร้างเอกสาร .docx ครบตามโครง OBE5 BRU (หมวด 1-6 + ลงนาม) คืนเป็น BytesIO - เรียก
    compute_offering_clo_achievement_raw() ครั้งเดียวแล้วส่งต่อให้ mco5_data_service.py จัดรูปแบบ (ข้อมูล
    ชุดเดียวกับ build_mco5_excel() เป๊ะ ไม่คำนวณซ้ำ) ไม่มีส่วนข้อมูลรายบุคคลตามสเปก"""
    raw_results = compute_offering_clo_achievement_raw(db, offering)
    clo_rows = compute_clo_rows(raw_results, target_rate)
    status_by_clo_id = {row.clo_id: row.status for row in clo_rows}
    course_info_rows = compute_course_info_rows(db, offering, target_rate)
    grade_dist = compute_grade_distribution(db, offering)
    assessment_rows = compute_assessment_confirmation_rows(db, offering, status_by_clo_id)

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = FONT_NAME
    style.font.size = BODY_SIZE
    # เผื่อ run ที่ไม่ได้ผ่าน _set_font เลย (เช่น run เริ่มต้นที่ python-docx อาจสร้างเองในบาง element)
    # ให้ตกไปใช้ style เริ่มต้นที่ตั้ง complex-script properties ไว้ถูกต้องแล้วเหมือนกัน ไม่ fallback
    # ไปฟอนต์/ขนาด default ของ Word เอง
    _apply_complex_script_properties(style.element.get_or_add_rPr(), BODY_SIZE)

    _add_header_obe5_label(doc)
    _add_title(doc)
    doc.add_paragraph()

    _add_category_heading(doc, "หมวดที่ 1 ข้อมูลทั่วไป")
    _add_course_info_table(doc, course_info_rows)

    _add_category_heading(doc, "หมวดที่ 2 การจัดการเรียนการสอนเปรียบเทียบกับแผนการสอน")
    _add_subitem(doc, "1. รายงานชั่วโมงการสอนจริงเทียบกับแผน")
    _add_placeholder(doc)
    _add_subitem(doc, "2. หัวข้อที่สอนไม่ครอบคลุมตามแผน (ถ้ามี) พร้อมเหตุผลและแนวทางแก้ไข")
    _add_placeholder(doc)
    _add_subitem(doc, "3. ประสิทธิผลของวิธีสอนที่ทำให้เกิดผลการเรียนรู้ตามผลลัพธ์การเรียนรู้ระดับรายวิชา (CLO)")
    _add_clo_table(doc, clo_rows)

    _add_category_heading(doc, "หมวดที่ 3 สรุปผลการจัดการเรียนการสอนของรายวิชา")
    _add_subitem(doc, "1. จำนวนนักศึกษาที่ลงทะเบียนเรียน")
    _add_paragraph(doc, f"{grade_dist.total_registered} คน")
    _add_subitem(doc, "2. จำนวนนักศึกษาที่คงอยู่เมื่อสิ้นสุดภาคการศึกษา")
    _add_paragraph(doc, f"{grade_dist.remaining} คน")
    _add_subitem(doc, "3. จำนวนนักศึกษาที่ถอน (W)")
    _add_paragraph(doc, f"{grade_dist.withdrawn} คน")
    _add_subitem(doc, "4. การกระจายเกรด")
    _add_grade_table(doc, grade_dist)
    _add_subitem(doc, "5. จำนวนนักศึกษาที่ได้รับผลกระทบจากการทุจริตทางวิชาการ (ถ้ามี)")
    _add_placeholder(doc)
    _add_subitem(doc, "6. ปัญหาของนักศึกษา/ทรัพยากรประกอบการเรียนการสอน")
    _add_placeholder(doc)
    _add_subitem(doc, "7. การประเมินผลลัพธ์การเรียนรู้ของรายวิชาจากผลการดำเนินการ")
    _add_assessment_table(doc, assessment_rows)

    _add_category_heading(doc, "หมวดที่ 4 ปัญหาและผลกระทบต่อการดำเนินการ")
    _add_placeholder(doc)

    _add_category_heading(doc, "หมวดที่ 5 การประเมินรายวิชา")
    _add_placeholder(doc)

    _add_category_heading(doc, "หมวดที่ 6 แผนการปรับปรุง")
    _add_placeholder(doc)

    _add_signature_section(doc, offering)

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
